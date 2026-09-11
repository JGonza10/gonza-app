"""
GONZA - Backend API (Flask + PostgreSQL)
Conecta la base de datos SQL con el frontend React.

CAMBIOS v2 (alineación con "Sistema de consulta de pagos"):
- Roles nuevos: consultor (solo lectura de prestamos/caja) y usuario (solo su info)
- CRUD completo de clientes y prestamos (PATCH/DELETE)
- Caja: incluye intereses (4%), fecha, nota, activo
- Ahorros y pagos a plazos: incluyen nota y activo
- Ruta /api/mis-datos para el rol "usuario"
"""

import os
import re
import secrets
import psycopg2
import psycopg2.extras
import itsdangerous
from datetime import date, datetime, time
from zoneinfo import ZoneInfo
from decimal import Decimal
from functools import wraps
from flask import Flask, jsonify, request, Response, g
from flask_cors import CORS
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from werkzeug.security import check_password_hash, generate_password_hash
import requests as http_requests
import pyotp
import qrcode
import io
import base64

app = Flask(__name__)
# En producción restringe a FRONTEND_URL (documentado en README/.env.example).
# Si no está configurada (desarrollo local), permite cualquier origen.
CORS(app, origins=os.environ.get("FRONTEND_URL") or "*")

# ─── ZONA HORARIA DEL NEGOCIO ─────────────────────────────────────────────────
# Todas las reglas de negocio con fecha (cortes de interés, "días para corte",
# vencidos, fecha por defecto de movimientos, timestamps de auditoría) deben
# calcularse en la hora de México, sin importar en qué zona horaria corra el
# servidor (Railway corre en UTC). get_db() fija esto mismo a nivel de sesión
# de PostgreSQL (ver "SET TIME ZONE" ahí) para que CURRENT_DATE/NOW() en SQL
# coincidan con lo que devuelven estas funciones.
ZONA_NEGOCIO_PG = "America/Mexico_City"
ZONA_NEGOCIO = ZoneInfo(ZONA_NEGOCIO_PG)


def ahora_mx():
    """Fecha y hora actuales en la zona horaria del negocio (Ciudad de México)."""
    return datetime.now(ZONA_NEGOCIO)


def hoy_mx():
    """Fecha de 'hoy' en la zona horaria del negocio (Ciudad de México).
       Usar SIEMPRE esta función (nunca date.today()/datetime.now() a secas)
       para cualquier cálculo de fecha que sea una regla de negocio."""
    return ahora_mx().date()

# ─── LÍMITE DE PETICIONES ─────────────────────────────────────────────────────
# Protección básica contra abuso/fuerza bruta a nivel de API completa. El login
# ya tiene su propio bloqueo por usuario (ver _cuenta_bloqueada); esto cubre el
# resto de rutas y pone un límite más estricto en las de acceso/recuperación.
# Nota: usa memoria en proceso (sin Redis), suficiente para un solo servidor
# pequeño; con varios workers cada uno lleva su propio conteo.
limiter = Limiter(
    app=app,
    key_func=get_remote_address,
    default_limits=["200 per hour"],
    storage_uri="memory://",
)


@app.route("/api/health", methods=["GET"])
@limiter.exempt
def health():
    """Endpoint de monitoreo (Railway/uptime). No requiere sesión ni cuenta
       para el límite de peticiones."""
    return jsonify({"status": "ok"})

# ─── TOKENS DE SESIÓN FIRMADOS ────────────────────────────────────────────────
# Reemplaza el header "X-Username" (que cualquiera podía falsificar) por un
# token firmado criptográficamente. Nadie puede fabricar uno válido sin conocer
# SECRET_KEY, que solo vive en el servidor.
#
# OBLIGATORIA: define SECRET_KEY como variable de entorno en Railway (o en tu
# .env local) con un valor largo y aleatorio, ej. generado con:
#   python -c "import secrets; print(secrets.token_hex(32))"
# Sin ella la app se niega a arrancar — nunca se usa una clave insegura por
# defecto ni se genera una aleatoria en cada arranque, porque eso invalidaría
# todas las sesiones activas cada vez que el proceso se reinicia.

SECRET_KEY = os.environ.get("SECRET_KEY")
if not SECRET_KEY:
    raise RuntimeError(
        "SECRET_KEY no está configurada. Defínela como variable de entorno "
        "antes de arrancar la aplicación (en Railway: pestaña Variables del "
        "servicio backend; en local: archivo .env). Genera un valor con: "
        'python -c "import secrets; print(secrets.token_hex(32))"'
    )

_serializer = itsdangerous.URLSafeTimedSerializer(SECRET_KEY)
TOKEN_SESION_MAX_AGE = 8 * 60 * 60   # 8 horas
TOKEN_RESET_MAX_AGE = 15 * 60        # 15 minutos


def _generar_token_sesion(username):
    return _serializer.dumps({"username": username}, salt="sesion")


def _obtener_username_autenticado():
    """Lee y valida el token del header Authorization: Bearer <token>.
       Devuelve (username, None) si es válido, o (None, mensaje_error) si no."""
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None, "No autorizado (falta sesión)"
    token = auth[7:].strip()
    try:
        datos = _serializer.loads(token, salt="sesion", max_age=TOKEN_SESION_MAX_AGE)
        return datos.get("username"), None
    except itsdangerous.SignatureExpired:
        return None, "Tu sesión expiró, inicia sesión de nuevo"
    except itsdangerous.BadSignature:
        return None, "Sesión inválida, inicia sesión de nuevo"

# ─── SERIALIZACIÓN GLOBAL DE FECHAS Y DECIMALES ──────────────────────────────
# Flask/json no sabe convertir objetos date/datetime/Decimal que vienen de
# PostgreSQL (vía RealDictCursor) a JSON. Sin esto, cualquier ruta que olvide
# castear una fecha con ::text revienta con un 500 silencioso en el navegador
# (pantalla en blanco). Este hook lo resuelve de forma global y definitiva.
_default_json = app.json.default

def _json_default(o):
    if isinstance(o, (date, datetime, time)):
        return o.isoformat()
    if isinstance(o, Decimal):
        return float(o)
    return _default_json(o)

app.json.default = _json_default


def _money(valor, default="0"):
    """Convierte un valor que llega del JSON de la petición (string, int,
       float o None) a Decimal seguro para dinero, redondeado a centavos.
       Usar SIEMPRE al leer montos de `data` antes de guardarlos — nunca
       float(), que puede arrastrar imprecisión binaria (float(1234.55) no
       siempre es exactamente 1234.55). Pasar por str() primero evita eso."""
    if valor is None or valor == "":
        valor = default
    return Decimal(str(valor)).quantize(Decimal("0.01"))

# ─── CONEXIÓN A LA BASE DE DATOS ─────────────────────────────────────────────
# DATABASE_URL viene de la variable de entorno (Railway la inyecta automático)
# En local, la defines en .env o en tu terminal

def get_db():
    """Devuelve una conexión a PostgreSQL.

    Fija la zona horaria de la SESIÓN a la del negocio (Ciudad de México),
    sin importar en qué zona corra el servidor (Railway corre en UTC). Así
    CURRENT_DATE, NOW() y CURRENT_TIMESTAMP —usados en cortes de interés,
    "días para corte", vencidos, y en los timestamps de auditoría/soft
    delete— reflejan el día calendario real de México, no el de UTC. Esto es
    lo correcto para reglas de negocio con dinero de por medio: usar el
    reloj de la computadora de quien esté usando el sistema en ese momento
    permitiría a cualquiera adelantar/atrasar su equipo para manipular
    cuándo se cobra un interés o cuándo algo se marca como vencido.
    """
    conn = psycopg2.connect(
        os.environ.get("DATABASE_URL"),
        cursor_factory=psycopg2.extras.RealDictCursor  # devuelve dicts, no tuplas
    )
    with conn.cursor() as cur:
        cur.execute("SET TIME ZONE %s;", (ZONA_NEGOCIO_PG,))
    return conn

# ─── RUTAS: AUTENTICACIÓN ─────────────────────────────────────────────────────

def _asegurar_tablas_auxiliares():
    """Crea tablas auxiliares del sistema si no existen todavía. Es seguro
       llamar esto en cada arranque: usa IF NOT EXISTS, nunca duplica ni
       borra nada. Así no dependemos de que alguien corra una migración
       manual — el propio despliegue se encarga."""
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS historial_accesos (
                id SERIAL PRIMARY KEY,
                username VARCHAR(100) NOT NULL,
                exito BOOLEAN NOT NULL,
                ip VARCHAR(64),
                fecha TIMESTAMP NOT NULL DEFAULT NOW()
            );
        """)
        cur.execute("""
            CREATE INDEX IF NOT EXISTS idx_historial_accesos_username_fecha
            ON historial_accesos (username, fecha DESC);
        """)

        # Borrado suave: las tablas financieras nunca se borran físicamente,
        # solo se marcan con eliminado_en. Así un borrado por error siempre
        # se puede recuperar, y queda registro de cuándo se "eliminó".
        for tabla in TABLAS_CON_BORRADO_SUAVE:
            cur.execute(f'ALTER TABLE {tabla} ADD COLUMN IF NOT EXISTS eliminado_en TIMESTAMP;')

        # Auditoría: quién creó/editó/eliminó qué registro y cuándo.
        cur.execute("""
            CREATE TABLE IF NOT EXISTS auditoria (
                id SERIAL PRIMARY KEY,
                tabla VARCHAR(50) NOT NULL,
                registro_id INTEGER NOT NULL,
                accion VARCHAR(20) NOT NULL,
                usuario VARCHAR(100),
                detalle JSONB,
                fecha TIMESTAMP NOT NULL DEFAULT NOW()
            );
        """)
        cur.execute("""
            CREATE INDEX IF NOT EXISTS idx_auditoria_tabla_registro
            ON auditoria (tabla, registro_id, fecha DESC);
        """)

        # 2FA (TOTP) para administradores: el secreto se guarda solo tras
        # confirmar un código válido (totp_habilitado pasa a TRUE en ese momento).
        cur.execute("ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS totp_secret VARCHAR(32);")
        cur.execute("ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS totp_habilitado BOOLEAN NOT NULL DEFAULT FALSE;")

        conn.commit()
        conn.close()
    except Exception as e:
        app.logger.warning(f"No se pudo verificar/crear tablas auxiliares: {e}")


# Tablas financieras con borrado suave (columna eliminado_en) en vez de DELETE físico.
TABLAS_CON_BORRADO_SUAVE = ("prestamos", "clientes", "ahorros", "caja", "pagos_plazos", "caja_movimientos")

_asegurar_tablas_auxiliares()


def _registrar_auditoria(cur, tabla, registro_id, accion, detalle=None):
    """Inserta un renglón de auditoría usando el cursor/transacción de la
       petición actual, para que quede atómico con el cambio que describe.
       accion: "crear" | "editar" | "eliminar"."""
    username, _ = _obtener_username_autenticado()
    cur.execute(
        "INSERT INTO auditoria (tabla, registro_id, accion, usuario, detalle) VALUES (%s, %s, %s, %s, %s);",
        (tabla, registro_id, accion, username, psycopg2.extras.Json(detalle) if detalle is not None else None),
    )

INTENTOS_MAXIMOS = 5
VENTANA_BLOQUEO_MINUTOS = 15


def _ip_del_cliente():
    """Railway pone la IP real del visitante en X-Forwarded-For (va detrás
       de un proxy), así que la revisamos primero."""
    xff = request.headers.get("X-Forwarded-For", "")
    if xff:
        return xff.split(",")[0].strip()
    return request.remote_addr or "desconocida"


def _registrar_acceso(username, exito):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO historial_accesos (username, exito, ip) VALUES (%s, %s, %s);",
        (username, exito, _ip_del_cliente())
    )
    conn.commit()
    conn.close()


def _cuenta_bloqueada(username):
    """True si hubo 5+ intentos fallidos en los últimos 15 minutos."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT COUNT(*) AS n FROM historial_accesos
        WHERE username = %s AND exito = FALSE
          AND fecha > NOW() - INTERVAL '%s minutes';
    """, (username, VENTANA_BLOQUEO_MINUTOS))
    n = cur.fetchone()["n"]
    conn.close()
    return n >= INTENTOS_MAXIMOS


@app.route("/api/login", methods=["POST"])
@limiter.limit("15 per minute")
def login():
    """Verifica usuario y contraseña, devuelve datos del usuario y su rol."""
    data = request.get_json()
    username = data.get("username", "")

    if username and _cuenta_bloqueada(username):
        return jsonify({
            "error": f"Demasiados intentos fallidos. Espera {VENTANA_BLOQUEO_MINUTOS} minutos antes de volver a intentar."
        }), 429

    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT u.id, u.username, u.nombre, u.password_hash, u.activo,
               u.cliente_id, r.nombre AS rol, u.totp_secret, u.totp_habilitado
        FROM usuarios u
        JOIN roles r ON u.rol_id = r.id
        WHERE u.username = %s;
    """, (username,))
    usuario = cur.fetchone()
    conn.close()

    if not usuario:
        _registrar_acceso(username, False)
        return jsonify({"error": "Usuario o contraseña incorrectos"}), 401

    if not usuario["activo"]:
        _registrar_acceso(username, False)
        return jsonify({"error": "Usuario inactivo"}), 403

    if not check_password_hash(usuario["password_hash"], data.get("password", "")):
        _registrar_acceso(username, False)
        return jsonify({"error": "Usuario o contraseña incorrectos"}), 401

    # Segundo factor (TOTP): si el usuario lo tiene activado, el código va en
    # el mismo POST (data.totp_code). Si falta o es incorrecto, se avisa con
    # requiere_2fa para que el frontend muestre ese campo sin pedir de nuevo
    # usuario/contraseña.
    if usuario["totp_habilitado"]:
        codigo = (data.get("totp_code") or "").strip()
        if not codigo or not pyotp.TOTP(usuario["totp_secret"]).verify(codigo, valid_window=1):
            _registrar_acceso(username, False)
            return jsonify({"error": "Código de verificación inválido o faltante", "requiere_2fa": True}), 401

    _registrar_acceso(username, True)
    return jsonify({
        "id": usuario["id"],
        "username": usuario["username"],
        "nombre": usuario["nombre"],
        "rol": usuario["rol"],
        "cliente_id": usuario["cliente_id"],
        "token": _generar_token_sesion(usuario["username"]),
    })

def get_usuario_actual(username):
    """Devuelve {rol, cliente_id} del usuario, o None si no existe/inactivo."""
    if not username:
        return None
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT r.nombre AS rol, u.cliente_id FROM usuarios u
        JOIN roles r ON u.rol_id = r.id
        WHERE u.username = %s AND u.activo = TRUE;
    """, (username,))
    row = cur.fetchone()
    conn.close()
    return dict(row) if row else None

def get_rol(username):
    """Devuelve el rol del usuario, o None si no existe (compatibilidad)."""
    info = get_usuario_actual(username)
    return info["rol"] if info else None


def requiere_rol(*roles_permitidos):
    """Decorador: bloquea la ruta si el token de sesión no es válido o el
       usuario no tiene uno de los roles permitidos."""
    def decorador(f):
        @wraps(f)
        def envoltura(*args, **kwargs):
            username, error = _obtener_username_autenticado()
            if error:
                return jsonify({"error": error}), 401
            rol = get_rol(username)
            if rol is None:
                return jsonify({"error": "No autorizado"}), 401
            if rol not in roles_permitidos:
                return jsonify({"error": "No tienes permiso para esta acción"}), 403
            g.username = username
            g.rol = rol
            return f(*args, **kwargs)
        return envoltura
    return decorador


def requiere_lectura(*roles_extra):
    """Decorador para rutas GET: permite administrador y analista siempre,
       y además los roles indicados en roles_extra (ej. 'consultor').
       Requiere sesión válida siempre (ya no permite acceso sin token)."""
    def decorador(f):
        @wraps(f)
        def envoltura(*args, **kwargs):
            username, error = _obtener_username_autenticado()
            if error:
                return jsonify({"error": error}), 401
            rol = get_rol(username)
            if rol is None:
                return jsonify({"error": "No autorizado"}), 401
            if rol not in ("administrador", "analista") + roles_extra:
                return jsonify({"error": "No tienes permiso para ver esta información"}), 403
            g.username = username
            g.rol = rol
            return f(*args, **kwargs)
        return envoltura
    return decorador


@app.route("/api/historial-accesos", methods=["GET"])
@requiere_rol("administrador")
def get_historial_accesos():
    """Últimos 200 accesos (exitosos y fallidos) — solo administrador."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT username, exito, ip, fecha::text AS fecha
        FROM historial_accesos
        ORDER BY fecha DESC
        LIMIT 200;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))


# ─── RUTAS: 2FA (TOTP) — solo administradores, sobre su propia cuenta ────────

@app.route("/api/2fa/estado", methods=["GET"])
@requiere_rol("administrador")
def get_2fa_estado():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT totp_habilitado FROM usuarios WHERE username = %s;", (g.username,))
    row = cur.fetchone()
    conn.close()
    return jsonify({"habilitado": bool(row and row["totp_habilitado"])})


@app.route("/api/2fa/generar", methods=["POST"])
@requiere_rol("administrador")
def generar_2fa():
    """Genera un secreto TOTP nuevo (aún no activo) y el QR para escanearlo
       con Google Authenticator / Authy. Queda pendiente hasta /2fa/activar."""
    secreto = pyotp.random_base32()
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "UPDATE usuarios SET totp_secret = %s, totp_habilitado = FALSE WHERE username = %s;",
        (secreto, g.username),
    )
    conn.commit()
    conn.close()

    uri = pyotp.TOTP(secreto).provisioning_uri(name=g.username, issuer_name="Sistema GONZA")
    qr = qrcode.make(uri)
    buffer = io.BytesIO()
    qr.save(buffer, format="PNG")
    qr_base64 = base64.b64encode(buffer.getvalue()).decode("ascii")

    return jsonify({"secreto": secreto, "qr": f"data:image/png;base64,{qr_base64}"})


@app.route("/api/2fa/activar", methods=["POST"])
@requiere_rol("administrador")
def activar_2fa():
    data = request.get_json()
    codigo = (data.get("codigo") or "").strip()

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT totp_secret FROM usuarios WHERE username = %s;", (g.username,))
    row = cur.fetchone()
    if not row or not row["totp_secret"]:
        conn.close()
        return jsonify({"error": "Primero genera un código QR con /2fa/generar"}), 400

    if not pyotp.TOTP(row["totp_secret"]).verify(codigo, valid_window=1):
        conn.close()
        return jsonify({"error": "Código incorrecto"}), 400

    cur.execute("UPDATE usuarios SET totp_habilitado = TRUE WHERE username = %s;", (g.username,))
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Verificación en dos pasos activada"})


@app.route("/api/2fa/desactivar", methods=["POST"])
@requiere_rol("administrador")
def desactivar_2fa():
    """Requiere la contraseña actual (no el código TOTP) para desactivar,
       por si el usuario perdió acceso a su app de autenticación."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT password_hash FROM usuarios WHERE username = %s;", (g.username,))
    row = cur.fetchone()
    if not row or not check_password_hash(row["password_hash"], data.get("password", "")):
        conn.close()
        return jsonify({"error": "Contraseña incorrecta"}), 401

    cur.execute(
        "UPDATE usuarios SET totp_habilitado = FALSE, totp_secret = NULL WHERE username = %s;",
        (g.username,),
    )
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Verificación en dos pasos desactivada"})


# ─── RUTAS: USUARIOS ──────────────────────────────────────────────────────────

@app.route("/api/usuarios", methods=["GET"])
@requiere_rol("administrador")
def get_usuarios():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT u.id, u.username, u.nombre, u.activo, u.correo, u.cliente_id,
               r.nombre AS rol, u.rol_id
        FROM usuarios u JOIN roles r ON u.rol_id = r.id
        ORDER BY u.id;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

@app.route("/api/usuarios", methods=["POST"])
@requiere_rol("administrador")
def add_usuario():
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT INTO usuarios (username, password_hash, nombre, rol_id, cliente_id, correo)
            VALUES (%s, %s, %s, %s, %s, %s)
            RETURNING id;
        """, (
            data["username"],
            generate_password_hash(data["password"], method="pbkdf2:sha256"),
            data["nombre"],
            data["rol_id"],
            data.get("cliente_id"),
            data.get("correo", ""),
        ))
        nuevo_id = cur.fetchone()["id"]
        conn.commit()
    except psycopg2.errors.UniqueViolation:
        conn.rollback()
        conn.close()
        return jsonify({"error": "Ese nombre de usuario ya existe"}), 400
    conn.close()
    return jsonify({"id": nuevo_id}), 201

@app.route("/api/usuarios/<int:uid>", methods=["PATCH"])
@requiere_rol("administrador")
def update_usuario(uid):
    """Activa/desactiva un usuario, cambia su rol, cliente vinculado o resetea su contraseña."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    if "activo" in data:
        cur.execute("UPDATE usuarios SET activo = %s WHERE id = %s;", (data["activo"], uid))
    if "rol_id" in data:
        cur.execute("UPDATE usuarios SET rol_id = %s WHERE id = %s;", (data["rol_id"], uid))
    if "cliente_id" in data:
        cur.execute("UPDATE usuarios SET cliente_id = %s WHERE id = %s;", (data["cliente_id"], uid))
    if "correo" in data:
        cur.execute("UPDATE usuarios SET correo = %s WHERE id = %s;", (data["correo"], uid))
    if "nombre" in data:
        cur.execute("UPDATE usuarios SET nombre = %s WHERE id = %s;", (data["nombre"], uid))
    if "password" in data and data["password"]:
        cur.execute("UPDATE usuarios SET password_hash = %s WHERE id = %s;",
                     (generate_password_hash(data["password"], method="pbkdf2:sha256"), uid))

    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Usuario actualizado"})

@app.route("/api/usuarios/<int:uid>", methods=["DELETE"])
@requiere_rol("administrador")
def delete_usuario(uid):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM usuarios WHERE id = %s;", (uid,))
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Usuario eliminado"})

@app.route("/api/roles", methods=["GET"])
@requiere_rol("administrador")
def get_roles():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, nombre FROM roles ORDER BY id;")
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))


# ─── RUTAS: PRÉSTAMOS ────────────────────────────────────────────────────────

@app.route("/api/buscar", methods=["GET"])
@requiere_lectura("consultor")
def buscar_global():
    """Busca un texto libre (nombre o teléfono) en clientes y en los
       nombres de deudor registrados en préstamos. Pensado para un buscador
       único arriba de todos los módulos."""
    q = (request.args.get("q") or "").strip()
    if len(q) < 2:
        return jsonify({"clientes": [], "prestamos": []})

    patron = f"%{q}%"
    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        SELECT id, nombre, apellido_pat, apellido_mat, telefono
        FROM clientes
        WHERE eliminado_en IS NULL
          AND (nombre ILIKE %s OR apellido_pat ILIKE %s OR apellido_mat ILIKE %s OR telefono ILIKE %s)
        ORDER BY nombre
        LIMIT 15;
    """, (patron, patron, patron, patron))
    clientes = cur.fetchall()

    cur.execute("""
        SELECT id, cliente_id, deudor_nombre, monto, pagado
        FROM prestamos
        WHERE eliminado_en IS NULL AND deudor_nombre ILIKE %s
        ORDER BY fecha_prestamo DESC
        LIMIT 15;
    """, (patron,))
    prestamos = cur.fetchall()

    conn.close()
    return jsonify({"clientes": list(clientes), "prestamos": list(prestamos)})


@app.route("/api/clientes/<int:cid>/historial-completo", methods=["GET"])
@requiere_lectura("consultor")
def get_historial_completo_cliente(cid):
    """Reúne en una sola respuesta todo lo relacionado a un cliente:
       datos básicos, préstamos, ahorros y aportaciones a caja."""
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT * FROM clientes WHERE id = %s;", (cid,))
    cliente = cur.fetchone()
    if not cliente:
        conn.close()
        return jsonify({"error": "Cliente no encontrado"}), 404

    cur.execute("""
        SELECT id, deudor_nombre, monto, interes_mensual, fecha_prestamo::text AS fecha_prestamo, pagado
        FROM prestamos WHERE cliente_id = %s AND eliminado_en IS NULL ORDER BY fecha_prestamo DESC;
    """, (cid,))
    prestamos = cur.fetchall()

    cur.execute("""
        SELECT id, cantidad, fecha::text AS fecha, nota
        FROM ahorros WHERE cliente_id = %s AND eliminado_en IS NULL ORDER BY fecha DESC;
    """, (cid,))
    ahorros = cur.fetchall()

    cur.execute("""
        SELECT c.id, c.fecha::text AS fecha, c.nota,
               COALESCE(SUM(cm.monto), 0) AS capital
        FROM caja c
        LEFT JOIN caja_movimientos cm ON cm.caja_id = c.id
        WHERE c.cliente_id = %s AND c.eliminado_en IS NULL
        GROUP BY c.id, c.fecha, c.nota
        ORDER BY c.fecha DESC;
    """, (cid,))
    caja = cur.fetchall()

    conn.close()
    return jsonify({
        "cliente": cliente,
        "prestamos": list(prestamos),
        "ahorros": list(ahorros),
        "caja": list(caja),
        "totales": {
            "prestado_activo": sum(float(p["monto"] or 0) for p in prestamos if not p["pagado"]),
            "ahorrado": sum(float(a["cantidad"] or 0) for a in ahorros),
            "en_caja": sum(float(c["capital"] or 0) for c in caja),
        },
    })


@app.route("/api/prestamos", methods=["GET"])
@requiere_lectura("consultor")
def get_prestamos():
    """Devuelve todos los préstamos (incluye fecha de último abono a capital)."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT p.*,
               (SELECT MAX(pp.fecha_pago) FROM pagos_prestamo pp
                 WHERE pp.prestamo_id = p.id AND pp.monto_capital > 0) AS fecha_abono_capital
        FROM prestamos p
        WHERE p.eliminado_en IS NULL
        ORDER BY p.fecha_prestamo ASC;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

@app.route("/api/prestamos", methods=["POST"])
@requiere_rol("administrador", "analista")
def add_prestamo():
    """Registra un nuevo préstamo. El deudor debe ser un cliente del catálogo."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    # Obtener el nombre completo del cliente seleccionado
    cur.execute("SELECT nombre, apellido_pat, apellido_mat FROM clientes WHERE id = %s;", (data["cliente_id"],))
    cliente = cur.fetchone()
    if not cliente:
        conn.close()
        return jsonify({"error": "Cliente no encontrado"}), 400

    deudor_nombre = f"{cliente['nombre']} {cliente['apellido_pat']} {cliente['apellido_mat'] or ''}".strip()

    cur.execute("""
        INSERT INTO prestamos (deudor_nombre, cliente_id, fecha_prestamo, monto, interes_mensual, nota)
        VALUES (%s, %s, %s, %s, %s, %s)
        RETURNING id;
    """, (
        deudor_nombre,
        data["cliente_id"],
        data["fecha_prestamo"],
        _money(data["monto"]),
        _money(data["interes_mensual"]),
        data.get("nota", "")
    ))
    nuevo_id = cur.fetchone()["id"]
    _registrar_auditoria(cur, "prestamos", nuevo_id, "crear", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id, "mensaje": "Préstamo registrado"}), 201

@app.route("/api/prestamos/<int:pid>", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def update_prestamo(pid):
    """Edita los datos generales de un préstamo (monto, interés, nota, fecha, estado activo)."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    campos_dinero = ("monto", "interes_mensual")
    campos = []
    valores = []
    for campo in ("fecha_prestamo", "monto", "interes_mensual", "nota", "activo"):
        if campo in data:
            campos.append(f"{campo} = %s")
            valores.append(_money(data[campo]) if campo in campos_dinero else data[campo])

    if not campos:
        conn.close()
        return jsonify({"error": "Nada para actualizar"}), 400

    valores.append(pid)
    cur.execute(f"UPDATE prestamos SET {', '.join(campos)} WHERE id = %s;", valores)
    _registrar_auditoria(cur, "prestamos", pid, "editar", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Préstamo actualizado"})

@app.route("/api/prestamos/<int:pid>", methods=["DELETE"])
@requiere_rol("administrador", "analista")
def delete_prestamo(pid):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE prestamos SET eliminado_en = NOW() WHERE id = %s;", (pid,))
    _registrar_auditoria(cur, "prestamos", pid, "eliminar")
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Préstamo eliminado"})

@app.route("/api/prestamos/<int:pid>/pagar", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def marcar_pagado(pid):
    """Marca un préstamo como pagado (liquidación total)."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        UPDATE prestamos
        SET pagado = TRUE,
            fecha_pago = %s,
            capital_abonado = monto,
            tipo_pago_id = (SELECT id FROM tipos_pago WHERE nombre = %s)
        WHERE id = %s;
    """, (data["fecha_pago"], data.get("tipo_pago", "transferencia"), pid))
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Marcado como pagado"})

@app.route("/api/prestamos/<int:pid>/abono", methods=["POST"])
@requiere_rol("administrador", "analista")
def registrar_abono(pid):
    """Registra un abono parcial (interés y/o capital) a un préstamo."""
    data = request.get_json()
    monto_interes = _money(data.get("monto_interes", 0))
    monto_capital = _money(data.get("monto_capital", 0))
    fecha_pago = data["fecha_pago"]
    tipo_pago = data.get("tipo_pago", "transferencia")
    nota = data.get("nota", "")

    if monto_interes <= 0 and monto_capital <= 0:
        return jsonify({"error": "Debes registrar al menos un monto mayor a 0"}), 400

    conn = get_db()
    cur = conn.cursor()

    # 1. Guardar el abono en el historial
    cur.execute("""
        INSERT INTO pagos_prestamo (prestamo_id, fecha_pago, monto_interes, monto_capital, tipo_pago_id, nota)
        VALUES (%s, %s, %s, %s, (SELECT id FROM tipos_pago WHERE nombre = %s), %s)
        RETURNING id;
    """, (pid, fecha_pago, monto_interes, monto_capital, tipo_pago, nota))
    nuevo_id = cur.fetchone()["id"]

    # 2. Actualizar el capital abonado del préstamo
    cur.execute("""
        UPDATE prestamos SET capital_abonado = capital_abonado + %s
        WHERE id = %s
        RETURNING monto, capital_abonado;
    """, (monto_capital, pid))
    row = cur.fetchone()

    # 3. Si el capital abonado ya cubre el monto total, marcar como pagado
    if row["capital_abonado"] >= row["monto"] and row["monto"] > 0:
        cur.execute("""
            UPDATE prestamos
            SET pagado = TRUE, fecha_pago = %s,
                tipo_pago_id = (SELECT id FROM tipos_pago WHERE nombre = %s)
            WHERE id = %s;
        """, (fecha_pago, tipo_pago, pid))

    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id, "mensaje": "Abono registrado"}), 201

@app.route("/api/prestamos/<int:pid>/historial", methods=["GET"])
@requiere_lectura("consultor")
def get_historial_prestamo(pid):
    """Devuelve el historial de abonos de un préstamo."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT pp.id, pp.fecha_pago::text AS fecha_pago, pp.monto_interes, pp.monto_capital, pp.nota, tp.nombre AS tipo_pago
        FROM pagos_prestamo pp
        LEFT JOIN tipos_pago tp ON pp.tipo_pago_id = tp.id
        WHERE pp.prestamo_id = %s
        ORDER BY pp.fecha_pago ASC, pp.id ASC;
    """, (pid,))
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

# ─── RUTAS: CORTES DE INTERÉS MENSUAL ────────────────────────────────────────

def _generar_cortes_faltantes(cur, pid, interes_mensual, fecha_prestamo):
    """
    Genera automáticamente los cortes mensuales que faltan para un préstamo,
    desde el mes siguiente a la fecha del préstamo hasta el mes actual.
    No crea cortes para meses futuros.
    """
    from dateutil.relativedelta import relativedelta

    hoy = hoy_mx()
    primer_corte = (fecha_prestamo + relativedelta(months=1)).replace(day=1)
    corte_actual = primer_corte

    while corte_actual <= hoy.replace(day=1):
        cur.execute("""
            INSERT INTO cortes_interes (prestamo_id, periodo, monto_interes)
            VALUES (%s, %s, %s)
            ON CONFLICT (prestamo_id, periodo) DO NOTHING;
        """, (pid, corte_actual, interes_mensual))
        corte_actual += relativedelta(months=1)


@app.route("/api/prestamos/<int:pid>/cortes", methods=["GET"])
@requiere_lectura("consultor")
def get_cortes_interes(pid):
    """
    Devuelve todos los cortes de interés de un préstamo.
    Genera automáticamente los cortes faltantes antes de responder.
    """
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT fecha_prestamo, interes_mensual, pagado FROM prestamos WHERE id = %s;", (pid,))
    prestamo = cur.fetchone()
    if not prestamo:
        conn.close()
        return jsonify({"error": "Préstamo no encontrado"}), 404

    if not prestamo["pagado"] and prestamo["interes_mensual"] > 0:
        _generar_cortes_faltantes(cur, pid, prestamo["interes_mensual"], prestamo["fecha_prestamo"])
        conn.commit()

    cur.execute("""
        SELECT ci.id,
               ci.periodo::text       AS periodo,
               ci.monto_interes,
               ci.pagado,
               ci.fecha_pago::text    AS fecha_pago,
               ci.monto_pagado,
               ci.nota,
               tp.nombre              AS tipo_pago
        FROM cortes_interes ci
        LEFT JOIN tipos_pago tp ON ci.tipo_pago_id = tp.id
        WHERE ci.prestamo_id = %s
        ORDER BY ci.periodo ASC;
    """, (pid,))
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))


@app.route("/api/prestamos/<int:pid>/cortes/<int:cid>/pagar", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def pagar_corte_interes(pid, cid):
    """
    Marca un corte de interés como pagado (total o parcial).
    Body: { fecha_pago, monto_pagado, tipo_pago, nota }
    """
    data = request.get_json()
    monto_pagado = _money(data.get("monto_pagado", 0))
    fecha_pago = data.get("fecha_pago")
    tipo_pago = data.get("tipo_pago", "transferencia")
    nota = data.get("nota", "")

    if monto_pagado <= 0:
        return jsonify({"error": "El monto pagado debe ser mayor a 0"}), 400

    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT monto_interes FROM cortes_interes WHERE id = %s AND prestamo_id = %s;", (cid, pid))
    corte = cur.fetchone()
    if not corte:
        conn.close()
        return jsonify({"error": "Corte no encontrado"}), 404

    monto_esperado = float(corte["monto_interes"])
    es_pagado_completo = monto_pagado >= monto_esperado

    cur.execute("""
        UPDATE cortes_interes
        SET pagado       = %s,
            fecha_pago   = %s,
            monto_pagado = monto_pagado + %s,
            tipo_pago_id = (SELECT id FROM tipos_pago WHERE nombre = %s),
            nota         = %s
        WHERE id = %s AND prestamo_id = %s;
    """, (es_pagado_completo, fecha_pago, monto_pagado, tipo_pago, nota, cid, pid))

    conn.commit()
    conn.close()
    return jsonify({
        "mensaje": "Corte marcado como pagado" if es_pagado_completo else "Abono parcial registrado",
        "pagado": es_pagado_completo,
    })


@app.route("/api/prestamos/<int:pid>/cortes/<int:cid>/prorrogar", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def prorrogar_corte(pid, cid):
    """Registra una prórroga: el interés no se cobró este mes."""
    data = request.get_json()
    nota = data.get("nota", "PRÓRROGA — interés no cobrado este mes")

    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        UPDATE cortes_interes
        SET nota = %s
        WHERE id = %s AND prestamo_id = %s;
    """, (nota, cid, pid))

    if cur.rowcount == 0:
        conn.close()
        return jsonify({"error": "Corte no encontrado"}), 404

    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Prórroga registrada"})


# ─── REFINANCIACIÓN / REESTRUCTURACIÓN DE PRÉSTAMOS ──────────────────────────

@app.route("/api/prestamos/refinanciar", methods=["POST"])
@requiere_rol("administrador", "analista")
def refinanciar_prestamos():
    """Consolida uno o varios préstamos activos de un mismo cliente en un
       préstamo nuevo: el saldo de capital pendiente (y, si se pide, el
       interés mensual vencido sin cobrar) de los préstamos elegidos se
       cierra como 'liquidado por refinanciación' y pasa a ser el monto
       del préstamo nuevo, con su propia fecha y tasa de interés."""
    data = request.get_json()
    cliente_id = data.get("cliente_id")
    try:
        prestamo_ids = list({int(x) for x in (data.get("prestamo_ids") or [])})
    except (TypeError, ValueError):
        return jsonify({"error": "Lista de préstamos inválida"}), 400
    fecha_prestamo = data.get("fecha_prestamo")
    interes_mensual = _money(data.get("interes_mensual", 0))
    incluir_intereses = bool(data.get("incluir_intereses_pendientes", True))
    nota_usuario = (data.get("nota") or "").strip()

    if not cliente_id or not prestamo_ids or not fecha_prestamo:
        return jsonify({"error": "Faltan cliente, préstamos a refinanciar o fecha"}), 400

    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT nombre, apellido_pat, apellido_mat FROM clientes WHERE id = %s AND eliminado_en IS NULL;", (cliente_id,))
    cliente = cur.fetchone()
    if not cliente:
        conn.close()
        return jsonify({"error": "Cliente no encontrado"}), 400
    deudor_nombre = f"{cliente['nombre']} {cliente['apellido_pat']} {cliente['apellido_mat'] or ''}".strip()

    cur.execute("""
        SELECT id, monto, capital_abonado, interes_mensual, fecha_prestamo
        FROM prestamos
        WHERE id = ANY(%s) AND cliente_id = %s AND pagado = FALSE AND eliminado_en IS NULL;
    """, (prestamo_ids, cliente_id))
    prestamos = cur.fetchall()
    if len(prestamos) != len(prestamo_ids):
        conn.close()
        return jsonify({"error": "Uno o más préstamos no son válidos para refinanciar (ya pagados, eliminados o de otro cliente)"}), 400

    total_saldo = sum((Decimal(str(p["monto"])) - Decimal(str(p["capital_abonado"]))) for p in prestamos)
    if total_saldo <= 0:
        conn.close()
        return jsonify({"error": "El saldo total de los préstamos seleccionados es cero"}), 400

    total_interes_pendiente = Decimal("0")
    if incluir_intereses:
        for p in prestamos:
            if p["interes_mensual"] and p["interes_mensual"] > 0:
                _generar_cortes_faltantes(cur, p["id"], p["interes_mensual"], p["fecha_prestamo"])
        cur.execute("""
            SELECT COALESCE(SUM(monto_interes), 0) AS total
            FROM cortes_interes WHERE prestamo_id = ANY(%s) AND pagado = FALSE;
        """, (prestamo_ids,))
        total_interes_pendiente = Decimal(str(cur.fetchone()["total"]))

    nuevo_monto = (total_saldo + total_interes_pendiente).quantize(Decimal("0.01"))

    nota_nuevo = f"Refinanciación de préstamo(s) #{', #'.join(str(p['id']) for p in prestamos)}."
    if nota_usuario:
        nota_nuevo += f" {nota_usuario}"

    cur.execute("""
        INSERT INTO prestamos (deudor_nombre, cliente_id, fecha_prestamo, monto, interes_mensual, nota)
        VALUES (%s, %s, %s, %s, %s, %s)
        RETURNING id;
    """, (deudor_nombre, cliente_id, fecha_prestamo, nuevo_monto, interes_mensual, nota_nuevo))
    nuevo_id = cur.fetchone()["id"]
    _registrar_auditoria(cur, "prestamos", nuevo_id, "crear",
                          detalle={"refinanciacion_de": prestamo_ids, "monto": str(nuevo_monto)})

    for p in prestamos:
        saldo_p = Decimal(str(p["monto"])) - Decimal(str(p["capital_abonado"]))
        if saldo_p > 0:
            cur.execute("""
                INSERT INTO pagos_prestamo (prestamo_id, fecha_pago, monto_interes, monto_capital, nota)
                VALUES (%s, %s, 0, %s, %s);
            """, (p["id"], fecha_prestamo, saldo_p, f"Cierre por refinanciación en préstamo #{nuevo_id}"))
        cur.execute("""
            UPDATE prestamos
            SET pagado = TRUE, fecha_pago = %s, capital_abonado = monto,
                nota = COALESCE(NULLIF(nota, ''), '') || %s
            WHERE id = %s;
        """, (fecha_prestamo, f" [Refinanciado en préstamo #{nuevo_id} el {fecha_prestamo}]", p["id"]))
        if incluir_intereses:
            cur.execute("""
                UPDATE cortes_interes
                SET pagado = TRUE, fecha_pago = %s, monto_pagado = monto_interes,
                    nota = COALESCE(NULLIF(nota, ''), '') || %s
                WHERE prestamo_id = %s AND pagado = FALSE;
            """, (fecha_prestamo, f" Capitalizado en préstamo #{nuevo_id}", p["id"]))
        _registrar_auditoria(cur, "prestamos", p["id"], "editar", detalle={"refinanciado_en": nuevo_id})

    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id, "monto": float(nuevo_monto), "mensaje": "Préstamos refinanciados"}), 201


@app.route("/api/intereses-pendientes", methods=["GET"])
@requiere_lectura("consultor")
def get_resumen_intereses_pendientes():
    """
    Resumen global: todos los préstamos activos con sus intereses pendientes.
    Primero genera cortes faltantes para todos los préstamos activos.
    """
    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        SELECT id, fecha_prestamo, interes_mensual
        FROM prestamos
        WHERE pagado = FALSE AND interes_mensual > 0 AND eliminado_en IS NULL;
    """)
    prestamos = cur.fetchall()
    for p in prestamos:
        _generar_cortes_faltantes(cur, p["id"], p["interes_mensual"], p["fecha_prestamo"])
    conn.commit()

    cur.execute("""
        SELECT
            p.id                        AS prestamo_id,
            p.cliente_id,
            p.deudor_nombre,
            p.monto,
            p.interes_mensual,
            p.fecha_prestamo::text      AS fecha_prestamo,
            COUNT(ci.id)                AS total_cortes,
            COUNT(ci.id) FILTER (WHERE ci.pagado = FALSE) AS cortes_pendientes,
            COALESCE(SUM(ci.monto_interes) FILTER (WHERE ci.pagado = FALSE), 0) AS total_interes_pendiente,
            COALESCE(SUM(ci.monto_pagado),  0)            AS total_interes_cobrado
        FROM prestamos p
        LEFT JOIN cortes_interes ci ON ci.prestamo_id = p.id
        WHERE p.pagado = FALSE AND p.monto > 0 AND p.eliminado_en IS NULL
        GROUP BY p.id, p.cliente_id, p.deudor_nombre, p.monto, p.interes_mensual, p.fecha_prestamo
        ORDER BY total_interes_pendiente DESC;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))


def _generar_xlsx_cartera_vencida():
    """Construye el reporte de cartera vencida: préstamos activos con al
       menos un corte de interés mensual pendiente de cobro."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font

    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        SELECT id, fecha_prestamo, interes_mensual
        FROM prestamos
        WHERE pagado = FALSE AND interes_mensual > 0 AND eliminado_en IS NULL;
    """)
    activos = cur.fetchall()
    for p in activos:
        _generar_cortes_faltantes(cur, p["id"], p["interes_mensual"], p["fecha_prestamo"])
    conn.commit()

    cur.execute("""
        SELECT
            p.id AS prestamo_id, p.deudor_nombre, c.telefono,
            p.monto, (p.monto - p.capital_abonado) AS saldo,
            p.interes_mensual, p.fecha_prestamo::text AS fecha_prestamo,
            COUNT(ci.id) FILTER (WHERE ci.pagado = FALSE)                       AS cortes_pendientes,
            COALESCE(SUM(ci.monto_interes) FILTER (WHERE ci.pagado = FALSE), 0) AS interes_pendiente,
            (MIN(ci.periodo) FILTER (WHERE ci.pagado = FALSE))::text            AS primer_corte_pendiente
        FROM prestamos p
        JOIN clientes c ON c.id = p.cliente_id
        LEFT JOIN cortes_interes ci ON ci.prestamo_id = p.id
        WHERE p.pagado = FALSE AND p.monto > 0 AND p.eliminado_en IS NULL
        GROUP BY p.id, p.deudor_nombre, c.telefono, p.monto, p.capital_abonado, p.interes_mensual, p.fecha_prestamo
        HAVING COUNT(ci.id) FILTER (WHERE ci.pagado = FALSE) > 0
        ORDER BY interes_pendiente DESC;
    """)
    filas = cur.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Cartera vencida"
    negrita = Font(bold=True)
    ws.append(["#", "Deudor", "Teléfono", "Monto prestado", "Saldo", "Interés mensual",
               "Meses de interés vencidos", "Interés pendiente", "Primer mes sin cobrar"])
    for celda in ws[1]:
        celda.font = negrita

    total_saldo = Decimal("0")
    total_interes = Decimal("0")
    for f in filas:
        ws.append([
            f["prestamo_id"], f["deudor_nombre"], f["telefono"] or "—",
            float(f["monto"] or 0), float(f["saldo"] or 0), float(f["interes_mensual"] or 0),
            f["cortes_pendientes"], float(f["interes_pendiente"] or 0),
            f["primer_corte_pendiente"] or "—",
        ])
        total_saldo += Decimal(str(f["saldo"] or 0))
        total_interes += Decimal(str(f["interes_pendiente"] or 0))

    ws.append([])
    ws.append(["", "TOTAL", "", "", float(total_saldo), "", "", float(total_interes), ""])
    for celda in ws[ws.max_row]:
        celda.font = negrita

    for col, ancho in zip("ABCDEFGHI", (6, 26, 14, 14, 14, 14, 12, 16, 18)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.route("/api/reportes/cartera-vencida/xlsx", methods=["GET"])
@requiere_lectura("consultor")
def get_cartera_vencida_xlsx():
    """Descarga en Excel los préstamos activos con intereses mensuales sin cobrar."""
    xlsx_bytes = _generar_xlsx_cartera_vencida()
    archivo = f"cartera_vencida_{hoy_mx().isoformat()}.xlsx"
    return Response(
        xlsx_bytes,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{archivo}"'},
    )


# ─── RUTAS: CLIENTES ──────────────────────────────────────────────────────────

@app.route("/api/clientes", methods=["GET"])
@requiere_lectura("consultor")
def get_clientes():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM clientes WHERE eliminado_en IS NULL ORDER BY apellido_pat, nombre;")
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

@app.route("/api/clientes", methods=["POST"])
@requiere_rol("administrador", "analista")
def add_cliente():
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO clientes (nombre, apellido_pat, apellido_mat, telefono, direccion)
        VALUES (%s, %s, %s, %s, %s)
        RETURNING id;
    """, (
        data["nombre"], data["apellido_pat"],
        data.get("apellido_mat", ""), data.get("telefono", ""), data.get("direccion", "")
    ))
    nuevo_id = cur.fetchone()["id"]
    _registrar_auditoria(cur, "clientes", nuevo_id, "crear", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id}), 201

@app.route("/api/clientes/<int:cid>", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def update_cliente(cid):
    """Edita los datos de un cliente, incluyendo activar/desactivar."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    campos = []
    valores = []
    for campo in ("nombre", "apellido_pat", "apellido_mat", "telefono", "direccion", "activo"):
        if campo in data:
            campos.append(f"{campo} = %s")
            valores.append(data[campo])

    if not campos:
        conn.close()
        return jsonify({"error": "Nada para actualizar"}), 400

    valores.append(cid)
    cur.execute(f"UPDATE clientes SET {', '.join(campos)} WHERE id = %s;", valores)
    _registrar_auditoria(cur, "clientes", cid, "editar", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Cliente actualizado"})

@app.route("/api/clientes/<int:cid>", methods=["DELETE"])
@requiere_rol("administrador", "analista")
def delete_cliente(cid):
    """Borrado suave: el cliente y su historial (préstamos, ahorros, caja)
       quedan intactos, solo se ocultan de las listas."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE clientes SET eliminado_en = NOW() WHERE id = %s;", (cid,))
    _registrar_auditoria(cur, "clientes", cid, "eliminar")
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Cliente eliminado"})

# ─── RUTAS: AHORROS ───────────────────────────────────────────────────────────

@app.route("/api/ahorros", methods=["GET"])
@requiere_lectura()
def get_ahorros():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT a.id, a.cliente_id, c.nombre, c.apellido_pat, c.apellido_mat,
               a.cantidad, a.fecha, a.nota, a.activo
        FROM ahorros a JOIN clientes c ON a.cliente_id = c.id
        WHERE a.eliminado_en IS NULL
        ORDER BY c.apellido_pat;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

@app.route("/api/ahorros", methods=["POST"])
@requiere_rol("administrador", "analista")
def add_ahorro():
    """Da de alta el registro de ahorro de un cliente (uno por cliente)."""
    data = request.get_json()

    if "cliente_id" not in data:
        return jsonify({"error": "Falta cliente_id"}), 400

    conn = get_db()
    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT INTO ahorros (cliente_id, cantidad, fecha, nota)
            VALUES (%s, %s, %s, %s)
            RETURNING id;
        """, (
            data["cliente_id"],
            _money(data.get("cantidad", 0)),
            data.get("fecha") or hoy_mx().isoformat(),
            data.get("nota", ""),
        ))
        nuevo_id = cur.fetchone()["id"]
        _registrar_auditoria(cur, "ahorros", nuevo_id, "crear", detalle=data)
        conn.commit()
    except psycopg2.errors.UniqueViolation:
        conn.rollback()
        conn.close()
        return jsonify({"error": "Este cliente ya tiene un registro de ahorro"}), 400
    except Exception as e:
        conn.rollback()
        conn.close()
        app.logger.error(f"No se pudo registrar el ahorro: {e}")
        return jsonify({"error": "No se pudo registrar el ahorro"}), 500
    conn.close()
    return jsonify({"id": nuevo_id}), 201

@app.route("/api/ahorros/<int:aid>", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def update_ahorro(aid):
    """Actualiza el monto, nota o estado de ahorro de un cliente."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    campos = ["updated_at = CURRENT_TIMESTAMP"]
    valores = []
    for campo in ("cantidad", "nota", "activo", "fecha"):
        if campo in data:
            campos.append(f"{campo} = %s")
            valores.append(_money(data[campo]) if campo == "cantidad" else data[campo])

    valores.append(aid)
    cur.execute(f"UPDATE ahorros SET {', '.join(campos)} WHERE id = %s;", valores)
    _registrar_auditoria(cur, "ahorros", aid, "editar", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Ahorro actualizado"})

@app.route("/api/ahorros/<int:aid>", methods=["DELETE"])
@requiere_rol("administrador", "analista")
def delete_ahorro(aid):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE ahorros SET eliminado_en = NOW() WHERE id = %s;", (aid,))
    _registrar_auditoria(cur, "ahorros", aid, "eliminar")
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Registro de ahorro eliminado"})

@app.route("/api/clientes-sin-ahorro", methods=["GET"])
@requiere_rol("administrador", "analista")
def get_clientes_sin_ahorro():
    """Clientes que aún no tienen registro en ahorros (para el formulario de alta)."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT c.id, c.nombre, c.apellido_pat, c.apellido_mat
        FROM clientes c
        WHERE c.eliminado_en IS NULL
          AND c.id NOT IN (SELECT cliente_id FROM ahorros WHERE eliminado_en IS NULL)
        ORDER BY c.apellido_pat;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

# ─── RUTAS: CAJA ──────────────────────────────────────────────────────────────

@app.route("/api/caja", methods=["GET"])
@requiere_lectura("consultor")
def get_caja():
    """
    Devuelve los participantes de la caja con el capital acumulado
    calculado en tiempo real como la suma de sus movimientos reales
    (no un campo fijo editable, para evitar desincronización).
    """
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT c.id, c.participante, c.cliente_id, c.cuota, c.fecha_inicio,
               c.fecha, c.nota, c.activo,
               COALESCE(SUM(cm.monto), 0) AS capital
        FROM caja c
        LEFT JOIN caja_movimientos cm ON cm.caja_id = c.id AND cm.eliminado_en IS NULL
        WHERE c.eliminado_en IS NULL
        GROUP BY c.id, c.participante, c.cliente_id, c.cuota, c.fecha_inicio,
                 c.fecha, c.nota, c.activo
        ORDER BY c.id;
    """)
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

@app.route("/api/caja", methods=["POST"])
@requiere_rol("administrador", "analista")
def add_caja():
    """Agrega un participante a la caja. El participante debe ser un cliente del catálogo."""
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT nombre, apellido_pat, apellido_mat FROM clientes WHERE id = %s;", (data["cliente_id"],))
    cliente = cur.fetchone()
    if not cliente:
        conn.close()
        return jsonify({"error": "Cliente no encontrado"}), 400

    participante = f"{cliente['nombre']} {cliente['apellido_pat']} {cliente['apellido_mat'] or ''}".strip()

    cur.execute("""
        INSERT INTO caja (participante, cliente_id, cuota, fecha_inicio, fecha, nota)
        VALUES (%s, %s, %s, %s, %s, %s)
        RETURNING id;
    """, (
        participante,
        data["cliente_id"],
        _money(data.get("cuota", 0)),
        data.get("fecha_inicio", ""),
        data.get("fecha"),
        data.get("nota", ""),
    ))
    nuevo_id = cur.fetchone()["id"]

    # Si se proporcionó un capital inicial, se registra como el primer movimiento real
    capital_inicial = _money(data.get("capital", 0))
    if capital_inicial > 0:
        cur.execute("""
            INSERT INTO caja_movimientos (caja_id, fecha, monto, nota)
            VALUES (%s, %s, %s, %s);
        """, (nuevo_id, data.get("fecha_inicio") or hoy_mx().isoformat(), capital_inicial, "Capital inicial"))

    _registrar_auditoria(cur, "caja", nuevo_id, "crear", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id}), 201

@app.route("/api/caja/<int:cid>", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def update_caja(cid):
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    campos = []
    valores = []
    for campo in ("participante", "cuota", "fecha_inicio", "fecha", "nota", "activo"):
        if campo in data:
            campos.append(f"{campo} = %s")
            valores.append(_money(data[campo]) if campo == "cuota" else data[campo])

    if not campos:
        conn.close()
        return jsonify({"error": "Nada para actualizar"}), 400

    valores.append(cid)
    cur.execute(f"UPDATE caja SET {', '.join(campos)} WHERE id = %s;", valores)
    _registrar_auditoria(cur, "caja", cid, "editar", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Actualizado"})

@app.route("/api/caja/<int:cid>", methods=["DELETE"])
@requiere_rol("administrador", "analista")
def delete_caja(cid):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE caja SET eliminado_en = NOW() WHERE id = %s;", (cid,))
    _registrar_auditoria(cur, "caja", cid, "eliminar")
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Eliminado"})

# ─── RUTAS: PAGOS A PLAZOS ────────────────────────────────────────────────────

@app.route("/api/plazos", methods=["GET"])
@requiere_lectura()
def get_plazos():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM pagos_plazos WHERE eliminado_en IS NULL ORDER BY id;")
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))

@app.route("/api/plazos", methods=["POST"])
@requiere_rol("administrador", "analista")
def add_plazo():
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO pagos_plazos (material, costo, meses_total, meses_pagados, cuota, abonado)
        VALUES (%s, %s, %s, %s, %s, %s)
        RETURNING id;
    """, (
        data["material"],
        _money(data["costo"]) if data.get("costo") is not None else None,
        data["meses_total"],
        data.get("meses_pagados", 0),
        _money(data["cuota"]) if data.get("cuota") is not None else None,
        _money(data.get("abonado", 0)),
    ))
    nuevo_id = cur.fetchone()["id"]
    _registrar_auditoria(cur, "pagos_plazos", nuevo_id, "crear", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id}), 201

@app.route("/api/plazos/<int:pid>", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def update_plazo(pid):
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()

    campos_dinero = ("costo", "cuota", "abonado")
    campos = []
    valores = []
    for campo in ("material", "costo", "meses_total", "meses_pagados", "cuota", "abonado", "nota", "fecha", "activo"):
        if campo in data:
            campos.append(f"{campo} = %s")
            if campo in campos_dinero:
                valores.append(_money(data[campo]) if data[campo] is not None else None)
            else:
                valores.append(data[campo])

    if not campos:
        conn.close()
        return jsonify({"error": "Nada para actualizar"}), 400

    valores.append(pid)
    cur.execute(f"UPDATE pagos_plazos SET {', '.join(campos)} WHERE id = %s;", valores)
    _registrar_auditoria(cur, "pagos_plazos", pid, "editar", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Actualizado"})

@app.route("/api/plazos/<int:pid>", methods=["DELETE"])
@requiere_rol("administrador", "analista")
def delete_plazo(pid):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE pagos_plazos SET eliminado_en = NOW() WHERE id = %s;", (pid,))
    _registrar_auditoria(cur, "pagos_plazos", pid, "eliminar")
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Eliminado"})

@app.route("/api/plazos/<int:pid>/abonar", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def abonar_plazo(pid):
    """Registra un abono mensual al artículo a plazos."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        UPDATE pagos_plazos
        SET meses_pagados = LEAST(meses_pagados + 1, meses_total),
            abonado = abonado + COALESCE(cuota, 0)
        WHERE id = %s
        RETURNING meses_pagados, abonado;
    """, (pid,))
    result = cur.fetchone()
    conn.commit()
    conn.close()
    return jsonify(dict(result))

# ─── RUTAS: MIS DATOS (rol "usuario") ────────────────────────────────────────

@app.route("/api/mis-datos", methods=["GET"])
@requiere_rol("usuario", "administrador", "analista", "consultor")
def get_mis_datos():
    """Para el rol 'usuario': devuelve préstamos, ahorro y caja del cliente
       vinculado a su cuenta. Otros roles pueden usarlo pasando ?cliente_id=."""
    username = g.username
    info = get_usuario_actual(username)

    cliente_id = request.args.get("cliente_id", type=int)
    if info["rol"] == "usuario":
        cliente_id = info["cliente_id"]

    if not cliente_id:
        return jsonify({"error": "No hay un cliente vinculado a este usuario"}), 400

    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT * FROM clientes WHERE id = %s;", (cliente_id,))
    cliente = cur.fetchone()

    cur.execute("""
        SELECT * FROM prestamos WHERE cliente_id = %s AND eliminado_en IS NULL ORDER BY fecha_prestamo ASC;
    """, (cliente_id,))
    prestamos = cur.fetchall()

    cur.execute("SELECT * FROM ahorros WHERE cliente_id = %s AND eliminado_en IS NULL;", (cliente_id,))
    ahorro = cur.fetchone()

    cur.execute("SELECT * FROM caja WHERE cliente_id = %s AND eliminado_en IS NULL;", (cliente_id,))
    caja = cur.fetchall()

    conn.close()
    return jsonify({
        "cliente": dict(cliente) if cliente else None,
        "prestamos": list(prestamos),
        "ahorro": dict(ahorro) if ahorro else None,
        "caja": list(caja),
    })

# ─── RUTAS: ALERTAS Y CONFIGURACIÓN ──────────────────────────────────────────

@app.route("/api/configuracion/dias_anticipacion", methods=["GET"])
@requiere_rol("administrador")
def get_dias_anticipacion():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    conn.close()
    return jsonify({"dias_anticipacion": int(row["valor"]) if row else 2})

@app.route("/api/configuracion/dias_anticipacion", methods=["PATCH"])
@requiere_rol("administrador")
def set_dias_anticipacion():
    data = request.get_json()
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        UPDATE configuracion SET valor = %s WHERE clave = 'dias_anticipacion_alerta';
    """, (str(int(data["dias_anticipacion"])),))
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Actualizado"})


# ─── BACKUP Y RESTAURACIÓN COMPLETA (ESTRUCTURA + DATOS) ─────────────────────
# Implementado 100% en Python con psycopg2 (sin depender de pg_dump/psql).
# Esto evita problemas de versión entre el cliente y el servidor de Postgres,
# y funciona sin importar qué builder use Railway (Nixpacks, Railpack, etc).

# Tipos de columna que necesitan longitud/precisión explícita en el CREATE TABLE
_TIPOS_CON_LONGITUD = {"character varying": "VARCHAR", "character": "CHAR"}
_TIPOS_CON_PRECISION = {"numeric": "NUMERIC"}
_TIPOS_SIMPLES = {
    "timestamp without time zone": "TIMESTAMP",
    "timestamp with time zone": "TIMESTAMPTZ",
    "double precision": "DOUBLE PRECISION",
    "boolean": "BOOLEAN",
    "text": "TEXT",
    "date": "DATE",
    "time without time zone": "TIME",
    "integer": "INTEGER",
    "bigint": "BIGINT",
    "smallint": "SMALLINT",
    "real": "REAL",
    "json": "JSON",
    "jsonb": "JSONB",
    "uuid": "UUID",
}


def _tipo_columna_sql(col):
    """Traduce un tipo de information_schema.columns a sintaxis SQL de CREATE TABLE."""
    tipo = col["data_type"]
    if tipo in _TIPOS_CON_LONGITUD:
        largo = col["character_maximum_length"]
        return f"{_TIPOS_CON_LONGITUD[tipo]}({largo})" if largo else _TIPOS_CON_LONGITUD[tipo]
    if tipo in _TIPOS_CON_PRECISION:
        precision, escala = col["numeric_precision"], col["numeric_scale"]
        return f"NUMERIC({precision},{escala or 0})" if precision else "NUMERIC"
    if tipo in _TIPOS_SIMPLES:
        return _TIPOS_SIMPLES[tipo]
    return (col.get("udt_name") or tipo).upper()


def _tablas_publicas(cur):
    """Lista las tablas base del esquema 'public', en orden alfabético."""
    cur.execute("""
        SELECT table_name FROM information_schema.tables
        WHERE table_schema = 'public' AND table_type = 'BASE TABLE'
        ORDER BY table_name;
    """)
    return [r["table_name"] for r in cur.fetchall()]


def _generar_create_table(cur, tabla):
    """Genera DROP + CREATE TABLE para una tabla, incluyendo llave primaria
       y detectando columnas SERIAL/BIGSERIAL a partir de su default (nextval)."""
    cur.execute("""
        SELECT column_name, data_type, udt_name, character_maximum_length,
               numeric_precision, numeric_scale, is_nullable, column_default
        FROM information_schema.columns
        WHERE table_schema = 'public' AND table_name = %s
        ORDER BY ordinal_position;
    """, (tabla,))
    columnas = cur.fetchall()

    lineas = []
    for c in columnas:
        es_serial = bool(c["column_default"]) and "nextval(" in c["column_default"]
        if es_serial:
            tipo_sql = "BIGSERIAL" if c["data_type"] == "bigint" else "SERIAL"
        else:
            tipo_sql = _tipo_columna_sql(c)

        linea = f'    "{c["column_name"]}" {tipo_sql}'
        if c["is_nullable"] == "NO":
            linea += " NOT NULL"
        if c["column_default"] is not None and not es_serial:
            linea += f' DEFAULT {c["column_default"]}'
        lineas.append(linea)

    cur.execute("""
        SELECT kcu.column_name
        FROM information_schema.table_constraints tc
        JOIN information_schema.key_column_usage kcu
          ON tc.constraint_name = kcu.constraint_name AND tc.table_schema = kcu.table_schema
        WHERE tc.table_schema = 'public' AND tc.table_name = %s
          AND tc.constraint_type = 'PRIMARY KEY'
        ORDER BY kcu.ordinal_position;
    """, (tabla,))
    pk_cols = [r["column_name"] for r in cur.fetchall()]
    if pk_cols:
        pk_list = ", ".join(f'"{c}"' for c in pk_cols)
        lineas.append(f"    PRIMARY KEY ({pk_list})")

    sql = f'DROP TABLE IF EXISTS "{tabla}" CASCADE;\n'
    sql += f'CREATE TABLE "{tabla}" (\n' + ",\n".join(lineas) + "\n);\n"
    return sql, [c["column_name"] for c in columnas]


def _generar_inserts(cur, tabla, columnas):
    """Genera los INSERT de todos los datos de una tabla, usando mogrify
       (adaptación segura de valores, igual que consultas parametrizadas)."""
    cur.execute(f'SELECT * FROM "{tabla}";')
    filas = cur.fetchall()
    if not filas:
        return ""

    cols_sql = ", ".join(f'"{c}"' for c in columnas)
    partes = [f'\n-- Datos: {tabla} ({len(filas)} filas)\n']
    for fila in filas:
        # Las columnas JSON/JSONB (ej. auditoria.detalle) llegan de psycopg2
        # ya decodificadas como dict/list de Python, y esos tipos no son
        # adaptables a SQL "en crudo" (igual que al insertarlas la primera
        # vez: _registrar_auditoria las envuelve con psycopg2.extras.Json).
        # Sin este envoltorio, mogrify truena con "can't adapt type 'dict'"
        # en cuanto la tabla auditoria tiene un solo registro, y el backup
        # completo fallaba con error 500.
        valores = tuple(
            psycopg2.extras.Json(fila[c]) if isinstance(fila[c], (dict, list)) else fila[c]
            for c in columnas
        )
        insert = cur.mogrify(
            f'INSERT INTO "{tabla}" ({cols_sql}) VALUES %s;', (valores,)
        )
        partes.append(insert.decode("utf-8") + "\n")
    return "".join(partes)


def _generar_setval_serial(cur, tabla, columnas):
    """Para columnas SERIAL, reajusta la secuencia al máximo valor insertado
       (si no se hace, el próximo INSERT normal choca con un ID duplicado)."""
    sql = ""
    for col in columnas:
        cur.execute("SELECT pg_get_serial_sequence(%s, %s);", (tabla, col))
        secuencia = cur.fetchone()["pg_get_serial_sequence"]
        if secuencia:
            sql += (
                f"SELECT setval('{secuencia}', "
                f'COALESCE((SELECT MAX("{col}") FROM "{tabla}"), 1), '
                f'(SELECT MAX("{col}") FROM "{tabla}") IS NOT NULL);\n'
            )
    return sql


def _generar_indices(cur, tabla):
    """Genera CREATE INDEX para los índices de una tabla (excluye el de la
       llave primaria, que ya se crea automáticamente con PRIMARY KEY)."""
    cur.execute("""
        SELECT indexname, indexdef FROM pg_indexes
        WHERE schemaname = 'public' AND tablename = %s
          AND indexname NOT LIKE '%%_pkey';
    """, (tabla,))
    sql = ""
    for idx in cur.fetchall():
        sql += f"{idx['indexdef']};\n"
    return sql


def _generar_vistas(cur):
    """Genera CREATE VIEW para las vistas del esquema public."""
    cur.execute("""
        SELECT viewname, definition FROM pg_views WHERE schemaname = 'public';
    """)
    sql = ""
    for v in cur.fetchall():
        sql += f'\nDROP VIEW IF EXISTS "{v["viewname"]}" CASCADE;\n'
        sql += f'CREATE VIEW "{v["viewname"]}" AS {v["definition"]}\n'
    return sql


@app.route("/api/configuracion/backup", methods=["GET"])
@requiere_rol("administrador")
def exportar_backup_completo():
    """Genera un dump completo (estructura + datos + vistas) en Python puro,
       sin depender de pg_dump, y lo entrega como .sql descargable."""
    conn = get_db()
    cur = conn.cursor()
    try:
        partes = ["-- Backup completo Sistema GONZA (estructura + datos)\n"
                   f"-- Generado: {ahora_mx().isoformat()}\n\n"
                   "BEGIN;\n"]

        tablas = _tablas_publicas(cur)
        columnas_por_tabla = {}
        for tabla in tablas:
            create_sql, columnas = _generar_create_table(cur, tabla)
            partes.append(f'\n-- ═══ Tabla: {tabla} ═══\n{create_sql}')
            columnas_por_tabla[tabla] = columnas

        for tabla in tablas:
            partes.append(_generar_inserts(cur, tabla, columnas_por_tabla[tabla]))

        partes.append("\n-- Índices\n")
        for tabla in tablas:
            partes.append(_generar_indices(cur, tabla))

        partes.append("\n-- Ajuste de secuencias (columnas SERIAL)\n")
        for tabla in tablas:
            partes.append(_generar_setval_serial(cur, tabla, columnas_por_tabla[tabla]))

        partes.append(_generar_vistas(cur))
        partes.append("\nCOMMIT;\n")
    except Exception as e:
        app.logger.error(f"Error al generar el backup: {e}")
        return jsonify({"error": "Error al generar el backup"}), 500
    finally:
        conn.close()

    contenido = "".join(partes)
    fecha = ahora_mx().strftime("%Y-%m-%d_%H-%M-%S")
    nombre_archivo = f"backup_gonza_{fecha}.sql"

    return Response(
        contenido,
        mimetype="application/sql",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


@app.route("/api/configuracion/restore", methods=["POST"])
@requiere_rol("administrador")
def restaurar_backup_completo():
    """Restaura la base de datos completa a partir de un .sql generado por
       /api/configuracion/backup. Se ejecuta como una sola transacción:
       si algo falla, no se aplica ningún cambio (rollback automático).
       ⚠️ SOBRESCRIBE los datos actuales de las tablas incluidas en el archivo."""
    if "archivo" not in request.files:
        return jsonify({"error": "No se envió ningún archivo"}), 400

    archivo = request.files["archivo"]
    if archivo.filename == "":
        return jsonify({"error": "Archivo vacío"}), 400
    if not archivo.filename.lower().endswith(".sql"):
        return jsonify({"error": "El archivo debe tener extensión .sql"}), 400

    contenido = archivo.read().decode("utf-8", errors="strict")
    if not contenido.strip():
        return jsonify({"error": "El archivo está vacío"}), 400

    conn = get_db()
    try:
        cur = conn.cursor()
        cur.execute(contenido)
        conn.commit()
    except Exception as e:
        conn.rollback()
        app.logger.error(f"Restauración de backup fallida: {e}")
        return jsonify({
            "error": "La restauración falló, no se aplicó ningún cambio (rollback automático). "
                     "Revisa los logs del servidor para más detalle."
        }), 500
    finally:
        conn.close()

    return jsonify({"mensaje": "Base de datos restaurada correctamente"})


@app.route("/api/dashboard/flujo-mensual", methods=["GET"])
@requiere_lectura("consultor")
def get_dashboard_flujo_mensual():
    """Ingresos cobrados (interés + abonos de capital) de los últimos 6 meses, mes a mes.
       Útil para ver la tendencia de flujo de caja en una gráfica."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        WITH meses AS (
            SELECT generate_series(
                date_trunc('month', CURRENT_DATE) - INTERVAL '5 months',
                date_trunc('month', CURRENT_DATE),
                INTERVAL '1 month'
            )::date AS mes
        ),
        interes AS (
            SELECT date_trunc('month', fecha_pago)::date AS mes, SUM(monto_pagado) AS total
            FROM cortes_interes
            WHERE pagado = TRUE AND fecha_pago >= date_trunc('month', CURRENT_DATE) - INTERVAL '5 months'
            GROUP BY 1
        ),
        capital AS (
            SELECT date_trunc('month', fecha_pago)::date AS mes, SUM(monto_capital) AS total
            FROM pagos_prestamo
            WHERE fecha_pago >= date_trunc('month', CURRENT_DATE) - INTERVAL '5 months'
            GROUP BY 1
        )
        SELECT meses.mes::text AS mes,
               COALESCE(interes.total, 0) AS interes_cobrado,
               COALESCE(capital.total, 0) AS capital_cobrado
        FROM meses
        LEFT JOIN interes ON interes.mes = meses.mes
        LEFT JOIN capital ON capital.mes = meses.mes
        ORDER BY meses.mes;
    """)
    rows = [dict(r) for r in cur.fetchall()]
    conn.close()
    return jsonify(rows)


@app.route("/api/dashboard/cartera", methods=["GET"])
@requiere_lectura("consultor")
def get_dashboard_cartera():
    """Clasifica los préstamos activos en vencidos / próximos a vencer / al día,
       según los días de anticipación configurados."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    dias_anticipacion = int(row["valor"]) if row else 2

    cur.execute("""
        SELECT v.id, v.deudor_nombre, v.monto, v.interes_mensual,
               v.proximo_corte::text AS proximo_corte,
               (v.proximo_corte - CURRENT_DATE) AS dias_para_corte
        FROM v_alertas_prestamos v
        JOIN prestamos p ON p.id = v.id
        WHERE p.eliminado_en IS NULL;
    """)
    rows = cur.fetchall()
    conn.close()

    vencidos = [r for r in rows if r["dias_para_corte"] < 0]
    proximos = [r for r in rows if 0 <= r["dias_para_corte"] <= dias_anticipacion]
    al_dia = [r for r in rows if r["dias_para_corte"] > dias_anticipacion]

    def _resumen(lista):
        return {"cantidad": len(lista), "monto": sum(float(r["monto"] or 0) for r in lista)}

    return jsonify({
        "vencidos": vencidos,
        "proximos": proximos,
        "al_dia": al_dia,
        "resumen": {"vencidos": _resumen(vencidos), "proximos": _resumen(proximos), "al_dia": _resumen(al_dia)},
    })


@app.route("/api/alertas", methods=["GET"])
@requiere_lectura("consultor", "usuario")
def get_alertas():
    """Préstamos cuyo próximo corte mensual está dentro de los días de anticipación configurados."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    dias = int(row["valor"]) if row else 2

    cur.execute("""
        SELECT v.id, v.deudor_nombre, v.monto, v.interes_mensual,
               v.fecha_base::text AS fecha_base,
               v.proximo_corte::text AS proximo_corte,
               (v.proximo_corte - CURRENT_DATE) AS dias_para_corte
        FROM v_alertas_prestamos v
        JOIN prestamos p ON p.id = v.id
        WHERE p.eliminado_en IS NULL AND (v.proximo_corte - CURRENT_DATE) BETWEEN 0 AND %s
        ORDER BY v.proximo_corte;
    """, (dias,))
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))


@app.route("/api/reportes/balance-caja", methods=["GET"])
@requiere_lectura("consultor")
def get_balance_caja_prestamos():
    """
    Compara, por persona, el saldo de capital que aún debe en préstamos
    activos contra el capital que ha aportado a la caja de ahorro.
    saldo = total_aportado - total_prestado
      positivo → aportó más de lo que debe
      negativo → debe más de lo que ha aportado
    Incluye el desglose de cada préstamo activo de la persona.
    """
    conn = get_db()
    cur = conn.cursor()

    # Préstamos activos (no pagados) por cliente, con su saldo de capital
    cur.execute("""
        SELECT p.id, p.cliente_id, p.deudor_nombre, p.fecha_prestamo::text,
               p.monto, p.capital_abonado,
               (p.monto - p.capital_abonado) AS saldo_capital
        FROM prestamos p
        WHERE p.eliminado_en IS NULL AND p.pagado = FALSE AND p.cliente_id IS NOT NULL
        ORDER BY p.fecha_prestamo ASC;
    """)
    prestamos = [dict(r) for r in cur.fetchall()]

    # Capital aportado a la caja por cliente (suma real de sus movimientos)
    cur.execute("""
        SELECT c.cliente_id, c.participante,
               COALESCE(SUM(cm.monto), 0) AS aportado
        FROM caja c
        LEFT JOIN caja_movimientos cm ON cm.caja_id = c.id AND cm.eliminado_en IS NULL
        WHERE c.eliminado_en IS NULL AND c.cliente_id IS NOT NULL
        GROUP BY c.cliente_id, c.participante;
    """)
    aportaciones = [dict(r) for r in cur.fetchall()]

    conn.close()

    personas = {}

    for p in prestamos:
        cid = p["cliente_id"]
        acc = personas.setdefault(cid, {
            "cliente_id": cid, "nombre": p["deudor_nombre"],
            "prestamos": [], "total_prestado": 0.0, "total_aportado": 0.0,
        })
        saldo_prestamo = float(p["saldo_capital"] or 0)
        acc["prestamos"].append({
            "id": p["id"],
            "fecha_prestamo": p["fecha_prestamo"],
            "monto": float(p["monto"] or 0),
            "saldo_capital": saldo_prestamo,
        })
        acc["total_prestado"] += saldo_prestamo

    for a in aportaciones:
        cid = a["cliente_id"]
        acc = personas.setdefault(cid, {
            "cliente_id": cid, "nombre": a["participante"],
            "prestamos": [], "total_prestado": 0.0, "total_aportado": 0.0,
        })
        acc["total_aportado"] += float(a["aportado"] or 0)
        if not acc["nombre"]:
            acc["nombre"] = a["participante"]

    resultado = []
    for acc in personas.values():
        acc["total_prestado"] = round(acc["total_prestado"], 2)
        acc["total_aportado"] = round(acc["total_aportado"], 2)
        acc["saldo"] = round(acc["total_aportado"] - acc["total_prestado"], 2)
        resultado.append(acc)

    resultado.sort(key=lambda x: (x["nombre"] or "").lower())

    totales = {
        "total_prestado": round(sum(p["total_prestado"] for p in resultado), 2),
        "total_aportado": round(sum(p["total_aportado"] for p in resultado), 2),
    }
    totales["saldo"] = round(totales["total_aportado"] - totales["total_prestado"], 2)

    return jsonify({"personas": resultado, "totales": totales})


# ─── ENVÍO DE CORREOS DE ALERTA ───────────────────────────────────────────────

def enviar_correo(destinatario, asunto, cuerpo_html, adjuntos=None):
    """Envía un correo usando la API de Resend.
       adjuntos: lista de dicts {"filename": str, "content": bytes (base64-able)}"""
    import base64
    api_key = os.environ.get("RESEND_API_KEY")
    payload = {
        "from": "GONZA <onboarding@resend.dev>",
        "to": [destinatario],
        "subject": asunto,
        "html": cuerpo_html,
    }
    if adjuntos:
        payload["attachments"] = [
            {"filename": a["filename"], "content": base64.b64encode(a["content"]).decode("ascii")}
            for a in adjuntos
        ]
    resp = http_requests.post(
        "https://api.resend.com/emails",
        headers={"Authorization": f"Bearer {api_key}"},
        json=payload,
        timeout=30,
    )
    if resp.status_code >= 400:
        raise Exception(f"Resend error {resp.status_code}: {resp.text}")


def get_destinatarios_admin():
    """Devuelve correos de administradores y analistas activos."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT u.correo, u.nombre FROM usuarios u
        JOIN roles r ON u.rol_id = r.id
        WHERE r.nombre IN ('administrador', 'analista') AND u.activo = TRUE
          AND u.correo IS NOT NULL AND u.correo != '';
    """)
    rows = cur.fetchall()
    conn.close()
    return list(rows)


# ─── RECORDATORIOS POR WHATSAPP (Twilio) ─────────────────────────────────────
# Requiere TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN y TWILIO_WHATSAPP_FROM en las
# variables de entorno (Railway). TWILIO_WHATSAPP_FROM es el número de Twilio
# con el prefijo "whatsapp:", ej. "whatsapp:+14155238886" (el del sandbox).
# Sin esas variables, el envío se omite y se reporta en la respuesta en vez
# de fallar silenciosamente — igual que RESEND_API_KEY/CRON_SECRET.

def _normalizar_telefono_mx(telefono):
    """Deja solo dígitos y antepone el código de país de México (52) si el
       número no lo trae ya. Devuelve None si no hay suficientes dígitos."""
    if not telefono:
        return None
    digitos = re.sub(r"\D", "", telefono)
    if len(digitos) == 10:
        digitos = "52" + digitos
    return f"+{digitos}" if len(digitos) >= 10 else None


def enviar_whatsapp(telefono, mensaje):
    """Envía un mensaje de WhatsApp vía la API de Twilio.
       Devuelve (True, None) si se envió, o (False, motivo) si no."""
    sid = os.environ.get("TWILIO_ACCOUNT_SID")
    token = os.environ.get("TWILIO_AUTH_TOKEN")
    remitente = os.environ.get("TWILIO_WHATSAPP_FROM")
    if not (sid and token and remitente):
        return False, "WhatsApp no configurado (falta TWILIO_ACCOUNT_SID/TWILIO_AUTH_TOKEN/TWILIO_WHATSAPP_FROM)"

    numero = _normalizar_telefono_mx(telefono)
    if not numero:
        return False, "Sin teléfono válido registrado"

    try:
        resp = http_requests.post(
            f"https://api.twilio.com/2010-04-01/Accounts/{sid}/Messages.json",
            auth=(sid, token),
            data={"From": remitente, "To": f"whatsapp:{numero}", "Body": mensaje},
            timeout=10,
        )
        if resp.status_code >= 300:
            return False, f"Twilio respondió {resp.status_code}: {resp.text[:200]}"
        return True, None
    except Exception as e:
        return False, str(e)


# ─── PROTECCIÓN CRON JOB ─────────────────────────────────────────────────────
# Las rutas de cron NO usan sesión de usuario sino un secret compartido.
# Define CRON_SECRET en las variables de entorno de Railway (o .env local).
# El cron job debe enviar el header:  X-Cron-Secret: <tu_secret>

def requiere_cron(f):
    """Decorador: permite la petición solo si el header X-Cron-Secret coincide
       con la variable de entorno CRON_SECRET.
       También acepta llamadas de un administrador autenticado (pruebas manuales).
       Si CRON_SECRET no está configurado, la ruta queda BLOQUEADA para peticiones
       no autenticadas (antes se dejaba pasar sin protección: cualquiera podía
       disparar envío de correos/WhatsApp/backups sin credenciales)."""
    @wraps(f)
    def envoltura(*args, **kwargs):
        cron_secret = os.environ.get("CRON_SECRET", "")

        # 1. Permitir si el header X-Cron-Secret es correcto (y está configurado)
        header_secret = request.headers.get("X-Cron-Secret", "")
        if cron_secret and header_secret == cron_secret:
            return f(*args, **kwargs)

        # 2. Permitir si es un administrador autenticado (prueba manual desde UI)
        username, error = _obtener_username_autenticado()
        if not error and username and get_rol(username) == "administrador":
            return f(*args, **kwargs)

        if not cron_secret:
            app.logger.warning(
                "CRON_SECRET no configurado: la ruta de cron permanece bloqueada "
                "para peticiones no autenticadas. Define esta variable de entorno "
                "en Railway para que el cron job externo pueda llamarla."
            )

        return jsonify({"error": "No autorizado. Se requiere X-Cron-Secret válido."}), 401
    return envoltura


@app.route("/api/alertas/enviar-correos", methods=["POST"])
@requiere_cron
def enviar_correos_alertas():
    """Envía un correo a cada usuario (admin/analista) con las alertas pendientes.
       Pensado para ser llamado por un cron job diario."""
    conn = get_db()
    cur = conn.cursor()

    # Días de anticipación configurados
    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    dias = int(row["valor"]) if row else 2

    # Alertas activas
    cur.execute("""
        SELECT v.deudor_nombre, v.monto, v.interes_mensual,
               v.proximo_corte::text AS proximo_corte,
               (v.proximo_corte - CURRENT_DATE) AS dias_para_corte
        FROM v_alertas_prestamos v
        JOIN prestamos p ON p.id = v.id
        WHERE p.eliminado_en IS NULL AND (v.proximo_corte - CURRENT_DATE) BETWEEN 0 AND %s
        ORDER BY v.proximo_corte;
    """, (dias,))
    alertas = cur.fetchall()
    conn.close()

    if not alertas:
        return jsonify({"mensaje": "Sin alertas pendientes, no se enviaron correos"})

    destinatarios = get_destinatarios_admin()
    if not destinatarios:
        return jsonify({"mensaje": "Sin destinatarios con correo configurado"}), 200

    # Construir el cuerpo del correo
    filas = "".join([
        f"<tr><td style='padding:6px 10px;border:1px solid #ddd'>{a['deudor_nombre']}</td>"
        f"<td style='padding:6px 10px;border:1px solid #ddd'>${a['interes_mensual']}</td>"
        f"<td style='padding:6px 10px;border:1px solid #ddd'>{a['proximo_corte']}</td>"
        f"<td style='padding:6px 10px;border:1px solid #ddd'>{'Hoy' if a['dias_para_corte']==0 else 'En ' + str(a['dias_para_corte']) + ' día(s)'}</td></tr>"
        for a in alertas
    ])
    cuerpo = f"""
    <h3>Sistema GONZA — Réditos próximos a vencer</h3>
    <table style="border-collapse:collapse;font-family:sans-serif;font-size:13px">
        <tr style="background:#0B1F4B;color:#C9A84C">
            <th style='padding:6px 10px;border:1px solid #ddd'>Deudor</th>
            <th style='padding:6px 10px;border:1px solid #ddd'>Interés</th>
            <th style='padding:6px 10px;border:1px solid #ddd'>Corte</th>
            <th style='padding:6px 10px;border:1px solid #ddd'>Vence</th>
        </tr>
        {filas}
    </table>
    """

    enviados = 0
    errores = []
    for d in destinatarios:
        try:
            enviar_correo(d["correo"], "GONZA — Alertas de réditos próximos a vencer", cuerpo)
            enviados += 1
        except Exception as e:
            errores.append(f"{d['correo']}: {str(e)}")

    return jsonify({"enviados": enviados, "errores": errores, "alertas": len(alertas)})


@app.route("/api/alertas/enviar-whatsapp", methods=["POST"])
@requiere_cron
def enviar_whatsapp_alertas():
    """Envía un recordatorio de WhatsApp directo a cada deudor con un corte
       de interés próximo a vencer (dentro de los días de anticipación
       configurados). A diferencia de /enviar-correos (que avisa al equipo),
       este le escribe al cliente mismo."""
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    dias = int(row["valor"]) if row else 2

    cur.execute("""
        SELECT p.deudor_nombre, p.interes_mensual, c.telefono,
               v.proximo_corte::text AS proximo_corte,
               (v.proximo_corte - CURRENT_DATE) AS dias_para_corte
        FROM v_alertas_prestamos v
        JOIN prestamos p ON p.id = v.id
        JOIN clientes c ON c.id = p.cliente_id
        WHERE p.eliminado_en IS NULL AND c.eliminado_en IS NULL
          AND (v.proximo_corte - CURRENT_DATE) BETWEEN 0 AND %s
        ORDER BY v.proximo_corte;
    """, (dias,))
    alertas = cur.fetchall()
    conn.close()

    if not alertas:
        return jsonify({"mensaje": "Sin alertas pendientes, no se enviaron mensajes"})

    enviados = 0
    errores = []
    for a in alertas:
        cuando = "hoy" if a["dias_para_corte"] == 0 else f"en {a['dias_para_corte']} día(s)"
        mensaje = (
            f"Hola {a['deudor_nombre']}, te recordamos que tu pago de interés de "
            f"${a['interes_mensual']} vence {cuando} ({a['proximo_corte']}). "
            "Gracias por tu puntualidad. — Sistema GONZA"
        )
        ok, motivo = enviar_whatsapp(a["telefono"], mensaje)
        if ok:
            enviados += 1
        else:
            errores.append(f"{a['deudor_nombre']}: {motivo}")

    return jsonify({"enviados": enviados, "errores": errores, "alertas": len(alertas)})


# ─── BACKUP DIARIO POR CORREO (CSV / XLSX / SQL) ─────────────────────────────

TABLAS_BACKUP = [
    "clientes", "prestamos", "pagos_prestamo", "cortes_interes",
    "ahorros", "caja", "caja_movimientos", "pagos_plazos",
    "tipos_pago", "configuracion", "usuarios", "roles",
]


def generar_backup_csv():
    """Devuelve una lista de adjuntos {filename, content} en CSV, uno por tabla."""
    import csv
    import io

    adjuntos = []
    conn = get_db()
    cur = conn.cursor()
    for tabla in TABLAS_BACKUP:
        cur.execute(f"SELECT * FROM {tabla};")
        rows = cur.fetchall()
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=rows[0].keys())
            writer.writeheader()
            for r in rows:
                writer.writerow(dict(r))
        adjuntos.append({"filename": f"{tabla}.csv", "content": buf.getvalue().encode("utf-8")})
    conn.close()
    return adjuntos


def generar_backup_xlsx():
    """Devuelve un único adjunto .xlsx con una hoja por tabla."""
    import io
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    conn = get_db()
    cur = conn.cursor()
    for tabla in TABLAS_BACKUP:
        cur.execute(f"SELECT * FROM {tabla};")
        rows = cur.fetchall()
        ws = wb.create_sheet(title=tabla[:31])
        if rows:
            headers = list(rows[0].keys())
            ws.append(headers)
            for r in rows:
                ws.append([str(v) if v is not None else "" for v in r.values()])
    conn.close()

    buf = io.BytesIO()
    wb.save(buf)
    return [{"filename": "gonza_backup.xlsx", "content": buf.getvalue()}]


def generar_backup_sql():
    """Genera un dump en formato SQL (INSERT statements) para cada tabla."""
    import io

    buf = io.StringIO()
    buf.write("-- GONZA backup generado automáticamente\n")

    conn = get_db()
    cur = conn.cursor()
    for tabla in TABLAS_BACKUP:
        cur.execute(f"SELECT * FROM {tabla};")
        rows = cur.fetchall()
        if not rows:
            continue
        cols = list(rows[0].keys())
        buf.write(f"\n-- Tabla: {tabla}\n")
        for r in rows:
            valores = []
            for c in cols:
                v = r[c]
                if v is None:
                    valores.append("NULL")
                elif isinstance(v, (int, float)):
                    valores.append(str(v))
                elif isinstance(v, bool):
                    valores.append("TRUE" if v else "FALSE")
                else:
                    valores.append("'" + str(v).replace("'", "''") + "'")
            buf.write(f"INSERT INTO {tabla} ({', '.join(cols)}) VALUES ({', '.join(valores)});\n")
    conn.close()

    return [{"filename": "gonza_backup.sql", "content": buf.getvalue().encode("utf-8")}]


@app.route("/api/backup/enviar", methods=["POST"])
@requiere_cron
def enviar_backup():
    """Genera un backup de la base de datos y lo envía por correo.
       Body opcional: {"formato": "csv" | "xlsx" | "sql"} (default: csv)
       Pensado para ser llamado por un cron job diario."""
    data = request.get_json(silent=True) or {}
    formato = data.get("formato", "csv")

    if formato == "csv":
        adjuntos = generar_backup_csv()
    elif formato == "xlsx":
        adjuntos = generar_backup_xlsx()
    elif formato == "sql":
        adjuntos = generar_backup_sql()
    else:
        return jsonify({"error": "Formato no soportado. Usa csv, xlsx o sql."}), 400

    destinatarios = get_destinatarios_admin()
    if not destinatarios:
        return jsonify({"mensaje": "Sin destinatarios con correo configurado"}), 200

    cuerpo = f"""
    <h3>Sistema GONZA — Respaldo diario de la base de datos</h3>
    <p>Se adjunta el respaldo en formato <b>{formato.upper()}</b> generado el día de hoy.</p>
    """

    enviados = 0
    errores = []
    for d in destinatarios:
        try:
            enviar_correo(d["correo"], f"GONZA — Respaldo diario ({formato.upper()})", cuerpo, adjuntos=adjuntos)
            enviados += 1
        except Exception as e:
            errores.append(f"{d['correo']}: {str(e)}")

    return jsonify({"enviados": enviados, "errores": errores, "formato": formato})


# ─── INFORME EJECUTIVO POR CORREO ────────────────────────────────────────────

@app.route("/api/resumen/informe", methods=["POST"])
@requiere_cron
def enviar_informe_resumen():
    """Genera un informe ejecutivo (resumen) en HTML y lo envía por correo.
       Pensado para ser llamado por un cron job diario o manual."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM v_resumen;")
    resumen = cur.fetchone()

    cur.execute("""
        SELECT deudor_nombre, SUM(monto - capital_abonado) AS total
        FROM prestamos WHERE pagado = FALSE AND monto > 0 AND eliminado_en IS NULL
        GROUP BY deudor_nombre ORDER BY total DESC LIMIT 5;
    """)
    top_deudores = cur.fetchall()
    conn.close()

    fmt = lambda v: f"${float(v or 0):,.2f}"

    filas_top = "".join([
        f"<tr><td style='padding:6px 10px;border:1px solid #ddd'>{d['deudor_nombre']}</td>"
        f"<td style='padding:6px 10px;border:1px solid #ddd'>{fmt(d['total'])}</td></tr>"
        for d in top_deudores
    ])

    cuerpo = f"""
    <h2 style="color:#0B1F4B">Sistema GONZA — Informe ejecutivo</h2>
    <table style="border-collapse:collapse;font-family:sans-serif;font-size:13px;margin-bottom:16px">
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Cartera activa</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt(resumen['cartera_activa'])}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Intereses por cobrar</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt(resumen.get('intereses_por_cobrar', 0))}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Capital caja de ahorro</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt(resumen['total_caja'])}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Intereses caja (4%)</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt(resumen.get('intereses_caja', 0))}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Total de ahorro</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt(resumen['total_ahorros'])}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Préstamos pendientes</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{resumen['prestamos_pendientes']}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Préstamos pagados</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{resumen['prestamos_cobrados']}</b></td></tr>
    </table>
    <h3 style="color:#0B1F4B">Top 5 deudores</h3>
    <table style="border-collapse:collapse;font-family:sans-serif;font-size:13px">
      <tr style="background:#0B1F4B;color:#C9A84C"><th style='padding:6px 10px;border:1px solid #ddd'>Deudor</th><th style='padding:6px 10px;border:1px solid #ddd'>Monto</th></tr>
      {filas_top}
    </table>
    """

    destinatarios = get_destinatarios_admin()
    if not destinatarios:
        return jsonify({"mensaje": "Sin destinatarios con correo configurado"}), 200

    enviados = 0
    errores = []
    for d in destinatarios:
        try:
            enviar_correo(d["correo"], "GONZA — Informe ejecutivo diario", cuerpo)
            enviados += 1
        except Exception as e:
            errores.append(f"{d['correo']}: {str(e)}")

    return jsonify({"enviados": enviados, "errores": errores})


# ─── RUTAS: MOVIMIENTOS DE CAJA ───────────────────────────────────────────────

@app.route("/api/caja/<int:cid>/movimientos", methods=["GET"])
@requiere_lectura("consultor")
def get_movimientos_caja(cid):
    """
    Devuelve el historial de aportaciones quincenales de un participante.
    Incluye el acumulado progresivo calculado con window function.
    """
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT
            id,
            fecha::text         AS fecha,
            monto,
            SUM(monto) OVER (
                ORDER BY fecha, id
                ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW
            )                   AS acumulado,
            nota
        FROM caja_movimientos
        WHERE caja_id = %s AND eliminado_en IS NULL
        ORDER BY fecha, id;
    """, (cid,))
    rows = cur.fetchall()
    conn.close()
    return jsonify(list(rows))


@app.route("/api/caja/<int:cid>/movimientos", methods=["POST"])
@requiere_rol("administrador", "analista")
def add_movimiento_caja(cid):
    """
    Registra una nueva aportación quincenal para un participante.
    Espera: { fecha, monto, nota }
    """
    data = request.get_json()

    monto = _money(data.get("monto", 0))
    if monto <= 0:
        return jsonify({"error": "El monto debe ser mayor a 0"}), 400

    conn = get_db()
    cur = conn.cursor()

    # Verificar que el participante existe
    cur.execute("SELECT id FROM caja WHERE id = %s;", (cid,))
    if not cur.fetchone():
        conn.close()
        return jsonify({"error": "Participante no encontrado"}), 404

    cur.execute("""
        INSERT INTO caja_movimientos (caja_id, fecha, monto, nota)
        VALUES (%s, %s, %s, %s)
        RETURNING id;
    """, (
        cid,
        data.get("fecha", "today"),
        monto,
        data.get("nota", ""),
    ))
    nuevo_id = cur.fetchone()["id"]
    _registrar_auditoria(cur, "caja_movimientos", nuevo_id, "crear", detalle=data)

    conn.commit()
    conn.close()
    return jsonify({"id": nuevo_id, "mensaje": "Aportación registrada"}), 201


@app.route("/api/caja/<int:cid>/movimientos/<int:mid>", methods=["PATCH"])
@requiere_rol("administrador", "analista")
def update_movimiento_caja(cid, mid):
    """Edita fecha, monto o nota de una aportación ya registrada. El capital se recalcula automáticamente."""
    data = request.get_json()

    campos = []
    valores = []
    for campo in ("fecha", "monto", "nota"):
        if campo in data:
            campos.append(f"{campo} = %s")
            valores.append(_money(data[campo]) if campo == "monto" else data[campo])

    if not campos:
        return jsonify({"error": "Nada para actualizar"}), 400

    conn = get_db()
    cur = conn.cursor()
    valores.extend([mid, cid])
    cur.execute(f"""
        UPDATE caja_movimientos SET {', '.join(campos)}
        WHERE id = %s AND caja_id = %s
        RETURNING id;
    """, valores)

    if not cur.fetchone():
        conn.close()
        return jsonify({"error": "Movimiento no encontrado"}), 404

    _registrar_auditoria(cur, "caja_movimientos", mid, "editar", detalle=data)
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Movimiento actualizado"})


@app.route("/api/caja/<int:cid>/movimientos/<int:mid>", methods=["DELETE"])
@requiere_rol("administrador")
def delete_movimiento_caja(cid, mid):
    """Elimina una aportación específica (solo administradores). El capital se recalcula automáticamente."""
    conn = get_db()
    cur = conn.cursor()

    cur.execute("UPDATE caja_movimientos SET eliminado_en = NOW() WHERE id = %s AND caja_id = %s RETURNING id;", (mid, cid))
    if not cur.fetchone():
        conn.close()
        return jsonify({"error": "Movimiento no encontrado"}), 404

    _registrar_auditoria(cur, "caja_movimientos", mid, "eliminar")
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Movimiento eliminado"})

# ─── RUTAS: RESET DE CONTRASEÑA (con código enviado al correo) ──────────────
# Flujo:
#   1. POST /api/usuarios/solicitar-reset {username}
#      -> genera un código de 6 dígitos, lo envía al correo registrado del
#         usuario, y devuelve un token firmado (válido 15 min) que contiene
#         el código esperado. El código NUNCA se guarda en la base de datos.
#   2. POST /api/usuarios/confirmar-reset {token, codigo, new_password}
#      -> valida que el token no haya expirado y que el código coincida:
#         solo entonces actualiza la contraseña.

@app.route("/api/usuarios/solicitar-reset", methods=["POST"])
@limiter.limit("5 per minute")
def solicitar_reset_password():
    data = request.get_json()
    username = (data.get("username") or "").strip()
    if not username:
        return jsonify({"error": "Username requerido"}), 400

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT nombre, correo, activo FROM usuarios WHERE username = %s;", (username,))
    row = cur.fetchone()
    conn.close()

    if not row:
        return jsonify({"error": "Usuario no encontrado"}), 404
    if not row["activo"]:
        return jsonify({"error": "Usuario inactivo"}), 403
    if not row["correo"]:
        return jsonify({
            "error": "Este usuario no tiene correo registrado. "
                     "Pide a un administrador que te agregue uno en el módulo de Usuarios."
        }), 400

    codigo = f"{secrets.randbelow(1_000_000):06d}"
    token = _serializer.dumps({"username": username, "codigo": codigo}, salt="reset-password")

    cuerpo = f"""
    <h3>Sistema GONZA — Código para restablecer tu contraseña</h3>
    <p>Hola {row['nombre']}, tu código de verificación es:</p>
    <p style="font-size:28px; font-weight:bold; letter-spacing:6px;">{codigo}</p>
    <p>Este código vence en 15 minutos. Si tú no solicitaste este cambio, ignora este correo.</p>
    """
    try:
        enviar_correo(row["correo"], "GONZA — Código para restablecer contraseña", cuerpo)
    except Exception as e:
        app.logger.error(f"No se pudo enviar el correo de reset a {row['correo']}: {e}")
        return jsonify({"error": "No se pudo enviar el correo, intenta de nuevo más tarde"}), 500

    return jsonify({
        "mensaje": f"Código enviado a {row['correo'][:3]}***",
        "token": token,
    })


@app.route("/api/usuarios/confirmar-reset", methods=["POST"])
@limiter.limit("10 per minute")
def confirmar_reset_password():
    data = request.get_json()
    token = data.get("token", "")
    codigo = (data.get("codigo") or "").strip()
    new_password = data.get("new_password", "")

    if not token or not codigo or not new_password:
        return jsonify({"error": "Faltan datos (token, código o nueva contraseña)"}), 400
    if len(new_password) < 6:
        return jsonify({"error": "La contraseña debe tener al menos 6 caracteres"}), 400

    try:
        datos = _serializer.loads(token, salt="reset-password", max_age=TOKEN_RESET_MAX_AGE)
    except itsdangerous.SignatureExpired:
        return jsonify({"error": "El código expiró, solicita uno nuevo"}), 400
    except itsdangerous.BadSignature:
        return jsonify({"error": "Token inválido, solicita un código nuevo"}), 400

    if not secrets.compare_digest(datos["codigo"], codigo):
        return jsonify({"error": "Código incorrecto"}), 400

    username = datos["username"]
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, activo FROM usuarios WHERE username = %s;", (username,))
    row = cur.fetchone()
    if not row or not row["activo"]:
        conn.close()
        return jsonify({"error": "Usuario no encontrado o inactivo"}), 404

    cur.execute(
        "UPDATE usuarios SET password_hash = %s WHERE username = %s;",
        (generate_password_hash(new_password, method="pbkdf2:sha256"), username)
    )
    conn.commit()
    conn.close()
    return jsonify({"mensaje": "Contraseña restablecida correctamente"})


# ─── RECORDATORIOS AUTOMÁTICOS (disparado por un cron externo, sin login) ─────
# No usa el sistema de tokens de sesión porque lo llama un servicio externo
# (Railway Cron, cron-job.org, etc.), no una persona con sesión iniciada.
# Protegido por CRON_SECRET: define esta variable de entorno en Railway con un
# valor largo y aleatorio (ej. python -c "import secrets; print(secrets.token_hex(32))"),
# y configura el cron para mandarlo en el header X-Cron-Secret.

@app.route("/api/cron/recordatorios", methods=["GET", "POST"])
def cron_recordatorios():
    secreto_esperado = os.environ.get("CRON_SECRET")
    secreto_recibido = request.headers.get("X-Cron-Secret", "")
    if not secreto_esperado or secreto_recibido != secreto_esperado:
        return jsonify({"error": "No autorizado"}), 403

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    dias = int(row["valor"]) if row else 2

    cur.execute("""
        SELECT v.deudor_nombre, v.monto, v.interes_mensual,
               v.proximo_corte::text AS proximo_corte,
               (v.proximo_corte - CURRENT_DATE) AS dias_para_corte
        FROM v_alertas_prestamos v
        JOIN prestamos p ON p.id = v.id
        WHERE p.eliminado_en IS NULL AND (v.proximo_corte - CURRENT_DATE) BETWEEN 0 AND %s
        ORDER BY v.proximo_corte;
    """, (dias,))
    alertas = cur.fetchall()
    conn.close()

    # Si no hay nada próximo a vencer, no manda correo (evita ruido diario innecesario)
    if not alertas:
        return jsonify({"enviado": False, "motivo": "Sin alertas pendientes hoy"})

    filas_lista = []
    for a in alertas:
        dias_texto = "Hoy" if a["dias_para_corte"] == 0 else f"En {a['dias_para_corte']} día(s)"
        filas_lista.append(
            f"<tr>"
            f"<td style='padding:6px 10px;border:1px solid #ddd'>{a['deudor_nombre']}</td>"
            f"<td style='padding:6px 10px;border:1px solid #ddd'>${float(a['interes_mensual'] or 0):,.2f}</td>"
            f"<td style='padding:6px 10px;border:1px solid #ddd'>{a['proximo_corte']}</td>"
            f"<td style='padding:6px 10px;border:1px solid #ddd'>{dias_texto}</td>"
            f"</tr>"
        )
    filas = "".join(filas_lista)
    cuerpo_html = f"""
    <div style="font-family:sans-serif;max-width:600px">
      <h2 style="color:#0B1F4B">JGM Gonzas Systems</h2>
      <h3 style="color:#0B1F4B;border-left:4px solid #C9A84C;padding-left:10px">
        🔔 Recordatorio automático: {len(alertas)} rédito(s) próximo(s) a vencer
      </h3>
      <table style="border-collapse:collapse;font-size:13px;width:100%">
        <tr style="background:#0B1F4B;color:#C9A84C">
          <th style="padding:6px 10px;border:1px solid #ddd;text-align:left">Deudor</th>
          <th style="padding:6px 10px;border:1px solid #ddd;text-align:left">Interés</th>
          <th style="padding:6px 10px;border:1px solid #ddd;text-align:left">Corte</th>
          <th style="padding:6px 10px;border:1px solid #ddd;text-align:left">Vence</th>
        </tr>
        {filas}
      </table>
      <p style="color:#888;font-size:11px;margin-top:20px">Correo automático diario de Sistema GONZA.</p>
    </div>
    """

    destinatarios = get_destinatarios_admin()
    enviados = []
    for d in destinatarios:
        try:
            enviar_correo(d["correo"], f"🔔 {len(alertas)} rédito(s) próximo(s) a vencer", cuerpo_html)
            enviados.append(d["correo"])
        except Exception as e:
            app.logger.warning(f"No se pudo enviar recordatorio a {d['correo']}: {e}")

    return jsonify({"enviado": True, "alertas": len(alertas), "destinatarios": enviados})




@app.route("/api/correo/completo", methods=["POST"])
@requiere_rol("administrador")
def enviar_correo_completo():
    """Genera y envía un único correo con 3 secciones:
       1. Alertas de réditos próximos
       2. Informe ejecutivo
       3. Respaldo de BD como adjunto Excel (.xlsx)
       Acepta body: {"formato": "xlsx" | "csv" | "sql"} (default: xlsx)
    """
    data = request.get_json(silent=True) or {}
    formato = data.get("formato", "xlsx")

    conn = get_db()
    cur = conn.cursor()

    # ── 1. Alertas ────────────────────────────────────────────────────────────
    cur.execute("SELECT valor FROM configuracion WHERE clave = 'dias_anticipacion_alerta';")
    row = cur.fetchone()
    dias = int(row["valor"]) if row else 2

    cur.execute("""
        SELECT v.deudor_nombre, v.monto, v.interes_mensual,
               v.proximo_corte::text AS proximo_corte,
               (v.proximo_corte - CURRENT_DATE) AS dias_para_corte
        FROM v_alertas_prestamos v
        JOIN prestamos p ON p.id = v.id
        WHERE p.eliminado_en IS NULL AND (v.proximo_corte - CURRENT_DATE) BETWEEN 0 AND %s
        ORDER BY v.proximo_corte;
    """, (dias,))
    alertas = cur.fetchall()

    if alertas:
        filas_alertas_lista = []
        for a in alertas:
            dias_texto = "Hoy" if a["dias_para_corte"] == 0 else f"En {a['dias_para_corte']} día(s)"
            interes_fmt = float(a["interes_mensual"] or 0)
            fila = (
                f"<tr>"
                f"<td style='padding:6px 10px;border:1px solid #ddd'>{a['deudor_nombre']}</td>"
                f"<td style='padding:6px 10px;border:1px solid #ddd'>${interes_fmt:,.2f}</td>"
                f"<td style='padding:6px 10px;border:1px solid #ddd'>{a['proximo_corte']}</td>"
                f"<td style='padding:6px 10px;border:1px solid #ddd'>{dias_texto}</td>"
                f"</tr>"
            )
            filas_alertas_lista.append(fila)
        filas_alertas = "".join(filas_alertas_lista)
        seccion_alertas = f"""
        <h3 style="color:#0B1F4B;border-left:4px solid #C9A84C;padding-left:10px">🔔 Alertas de réditos próximos ({len(alertas)})</h3>
        <table style="border-collapse:collapse;font-family:sans-serif;font-size:13px;margin-bottom:24px;width:100%">
          <tr style="background:#0B1F4B;color:#C9A84C">
            <th style='padding:6px 10px;border:1px solid #ddd;text-align:left'>Deudor</th>
            <th style='padding:6px 10px;border:1px solid #ddd;text-align:left'>Interés</th>
            <th style='padding:6px 10px;border:1px solid #ddd;text-align:left'>Corte</th>
            <th style='padding:6px 10px;border:1px solid #ddd;text-align:left'>Vence</th>
          </tr>
          {filas_alertas}
        </table>
        """
    else:
        seccion_alertas = """
        <h3 style="color:#0B1F4B;border-left:4px solid #C9A84C;padding-left:10px">🔔 Alertas de réditos</h3>
        <p style="color:#555;font-size:13px">✅ Sin alertas pendientes en los próximos días.</p>
        """

    # ── 2. Informe ejecutivo ──────────────────────────────────────────────────
    cur.execute("SELECT * FROM v_resumen;")
    resumen = cur.fetchone()

    cur.execute("""
        SELECT deudor_nombre, SUM(monto - capital_abonado) AS total
        FROM prestamos WHERE pagado = FALSE AND monto > 0 AND eliminado_en IS NULL
        GROUP BY deudor_nombre ORDER BY total DESC LIMIT 5;
    """)
    top_deudores = cur.fetchall()
    conn.close()

    fmt_m = lambda v: f"${float(v or 0):,.2f}"

    filas_top = "".join([
        f"<tr>"
        f"<td style='padding:6px 10px;border:1px solid #ddd'>{d['deudor_nombre']}</td>"
        f"<td style='padding:6px 10px;border:1px solid #ddd'>{fmt_m(d['total'])}</td>"
        f"</tr>"
        for d in top_deudores
    ])

    seccion_informe = f"""
    <h3 style="color:#0B1F4B;border-left:4px solid #C9A84C;padding-left:10px">📊 Informe ejecutivo</h3>
    <table style="border-collapse:collapse;font-family:sans-serif;font-size:13px;margin-bottom:16px">
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Cartera activa</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt_m(resumen['cartera_activa'])}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Capital caja de ahorro</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt_m(resumen['total_caja'])}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Total de ahorro</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{fmt_m(resumen['total_ahorros'])}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Préstamos pendientes</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{resumen['prestamos_pendientes']}</b></td></tr>
      <tr><td style='padding:6px 10px;border:1px solid #ddd'>Préstamos pagados</td><td style='padding:6px 10px;border:1px solid #ddd'><b>{resumen['prestamos_cobrados']}</b></td></tr>
    </table>
    <h4 style="color:#0B1F4B">Top 5 deudores</h4>
    <table style="border-collapse:collapse;font-family:sans-serif;font-size:13px;margin-bottom:24px">
      <tr style="background:#0B1F4B;color:#C9A84C">
        <th style='padding:6px 10px;border:1px solid #ddd;text-align:left'>Deudor</th>
        <th style='padding:6px 10px;border:1px solid #ddd;text-align:left'>Monto</th>
      </tr>
      {filas_top}
    </table>
    """

    # ── 3. Adjunto: respaldo ──────────────────────────────────────────────────
    if formato == "csv":
        adjuntos = generar_backup_csv()
    elif formato == "sql":
        adjuntos = generar_backup_sql()
    else:
        adjuntos = generar_backup_xlsx()
        formato = "xlsx"

    seccion_backup = f"""
    <h3 style="color:#0B1F4B;border-left:4px solid #C9A84C;padding-left:10px">💾 Respaldo de base de datos</h3>
    <p style="font-family:sans-serif;font-size:13px;color:#555">
      Se adjunta el respaldo en formato <b>{formato.upper()}</b> generado hoy.
    </p>
    """

    # ── Cuerpo final ──────────────────────────────────────────────────────────
    cuerpo = f"""
    <div style="font-family:sans-serif;max-width:680px;margin:0 auto">
      <div style="background:#0B1F4B;padding:18px 24px;border-radius:8px 8px 0 0">
        <h2 style="color:#C9A84C;margin:0">Sistema GONZA</h2>
        <p style="color:#8fa8c8;margin:4px 0 0;font-size:12px">Informe diario automático — {hoy_mx().strftime('%d/%m/%Y')}</p>
      </div>
      <div style="background:#fff;padding:24px;border:1px solid #ddd;border-top:none;border-radius:0 0 8px 8px">
        {seccion_alertas}
        {seccion_informe}
        {seccion_backup}
        <hr style="border:none;border-top:1px solid #eee;margin:20px 0"/>
        <p style="font-size:11px;color:#aaa">Este correo fue generado automáticamente por el Sistema GONZA.</p>
      </div>
    </div>
    """

    destinatarios = get_destinatarios_admin()
    if not destinatarios:
        return jsonify({"mensaje": "Sin destinatarios con correo configurado"}), 200

    enviados = 0
    errores = []
    for d in destinatarios:
        try:
            enviar_correo(
                d["correo"],
                f"GONZA — Informe diario {hoy_mx().strftime('%d/%m/%Y')}",
                cuerpo,
                adjuntos=adjuntos
            )
            enviados += 1
        except Exception as e:
            errores.append(f"{d['correo']}: {str(e)}")

    return jsonify({
        "enviados": enviados,
        "errores": errores,
        "alertas": len(alertas),
        "formato": formato
    })


# ─── INFORME POR DEUDOR ───────────────────────────────────────────────────────

def _datos_informe_deudor(cliente_id):
    """Consulta y arma los datos del informe de un deudor a partir de su
       cliente_id (NO de deudor_nombre: ese campo es una foto del nombre al
       momento de crear cada préstamo, así que si el cliente se editó entre
       un préstamo y otro, dos préstamos del mismo deudor pueden tener
       cadenas distintas y un match por texto se perdía alguno).
       Reutilizada por la vista JSON (para la pantalla) y la vista PDF (para descargar)."""
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT nombre, apellido_pat, apellido_mat FROM clientes WHERE id = %s;", (cliente_id,))
    cliente = cur.fetchone()
    nombre_actual = (
        f"{cliente['nombre']} {cliente['apellido_pat']} {cliente['apellido_mat'] or ''}".strip()
        if cliente else "Cliente no encontrado"
    )

    # Todos los préstamos del deudor
    cur.execute("""
        SELECT p.id, p.deudor_nombre, p.fecha_prestamo::text, p.monto,
               p.interes_mensual, p.pagado, p.fecha_pago::text, p.nota,
               p.capital_abonado,
               (p.monto - p.capital_abonado) AS saldo_capital,
               tp.nombre AS tipo_pago
        FROM prestamos p
        LEFT JOIN tipos_pago tp ON p.tipo_pago_id = tp.id
        WHERE p.cliente_id = %s AND p.eliminado_en IS NULL
        ORDER BY p.fecha_prestamo DESC;
    """, (cliente_id,))
    prestamos = [dict(r) for r in cur.fetchall()]

    ids = [p["id"] for p in prestamos]

    # Cortes de interés de todos sus préstamos
    cortes = []
    if ids:
        cur.execute("""
            SELECT ci.prestamo_id, ci.periodo::text, ci.monto_interes,
                   ci.pagado, ci.fecha_pago::text, ci.monto_pagado, ci.nota,
                   tp.nombre AS tipo_pago
            FROM cortes_interes ci
            LEFT JOIN tipos_pago tp ON ci.tipo_pago_id = tp.id
            WHERE ci.prestamo_id = ANY(%s)
            ORDER BY ci.periodo DESC;
        """, (ids,))
        cortes = [dict(r) for r in cur.fetchall()]

    # Abonos de capital
    abonos = []
    if ids:
        cur.execute("""
            SELECT pp.prestamo_id, pp.fecha_pago::text, pp.monto_interes,
                   pp.monto_capital, pp.nota, tp.nombre AS tipo_pago
            FROM pagos_prestamo pp
            LEFT JOIN tipos_pago tp ON pp.tipo_pago_id = tp.id
            WHERE pp.prestamo_id = ANY(%s)
            ORDER BY pp.fecha_pago DESC;
        """, (ids,))
        abonos = [dict(r) for r in cur.fetchall()]

    conn.close()

    # Totales
    activos = [p for p in prestamos if not p["pagado"] and p["monto"] > 0]
    # Saldo pendiente real (monto - capital_abonado), no el monto original:
    # si no se resta lo ya abonado, el informe muestra como "prestado" capital
    # que el deudor ya pagó.
    total_prestado    = sum(float(p["saldo_capital"] or 0) for p in activos)
    total_interes_mes = sum(float(p["interes_mensual"] or 0) for p in activos)
    total_pendiente   = sum(float(c["monto_interes"] or 0) for c in cortes if not c["pagado"])
    total_cobrado     = sum(float(c["monto_pagado"] or 0) for c in cortes)

    return {
        "deudor": nombre_actual,
        "prestamos": prestamos,
        "cortes": cortes,
        "abonos": abonos,
        "resumen": {
            "total_prestado": total_prestado,
            "total_interes_mensual": total_interes_mes,
            "interes_pendiente_acumulado": total_pendiente,
            "interes_cobrado_total": total_cobrado,
            "prestamos_activos": len(activos),
            "prestamos_pagados": len([p for p in prestamos if p["pagado"]]),
        }
    }


@app.route("/api/clientes/<int:cliente_id>/informe-deudor", methods=["GET"])
@requiere_lectura("consultor")
def get_informe_deudor(cliente_id):
    """Informe ejecutivo completo de un deudor: préstamos, abonos, cortes de interés."""
    return jsonify(_datos_informe_deudor(cliente_id))


def _generar_pdf_informe(datos):
    """Construye el PDF del informe ejecutivo a partir de los datos de _datos_informe_deudor."""
    from io import BytesIO
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    NAVY = colors.HexColor("#0B1F4B")
    GOLD = colors.HexColor("#C9A84C")
    ORANGE = colors.HexColor("#E87722")
    GRAY = colors.HexColor("#3B3B4F")
    BORDE = colors.HexColor("#D4D4DC")

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=16 * mm, bottomMargin=15 * mm,
                             leftMargin=18 * mm, rightMargin=18 * mm)
    styles = getSampleStyleSheet()
    titulo_style = ParagraphStyle("titulo", parent=styles["Title"], textColor=NAVY, fontSize=17, spaceAfter=1, alignment=0)
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], textColor=GOLD, fontSize=10.5, spaceAfter=0, alignment=0)
    seccion_style = ParagraphStyle("seccion", parent=styles["Heading2"], textColor=NAVY, fontSize=13,
                                    spaceBefore=14, spaceAfter=6)

    # Membrete: logo a la izquierda + nombre/subtítulo a la derecha (si el logo existe en el deploy)
    logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "gonza-systems.png")
    texto_membrete = [Paragraph("JGM Gonzas Systems", titulo_style), Paragraph("Informe ejecutivo de deudor", sub_style)]
    if os.path.exists(logo_path):
        logo_img = Image(logo_path, width=15 * mm, height=15 * mm)
        membrete = Table([[logo_img, texto_membrete]], colWidths=[20 * mm, 150 * mm])
        membrete.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (0, 0), 0)]))
        el = [membrete, Spacer(1, 10)]
    else:
        el = [Paragraph("JGM Gonzas Systems", titulo_style), Paragraph("Informe ejecutivo de deudor", sub_style), Spacer(1, 10)]

    el += [
        Paragraph(f"<b>Deudor:</b> {datos['deudor']}", styles["Normal"]),
        Paragraph(f"<b>Fecha del informe:</b> {hoy_mx().strftime('%d/%m/%Y')}", styles["Normal"]),
        Spacer(1, 10),
    ]

    r = datos["resumen"]
    resumen_data = [
        ["Capital prestado activo", f"${r['total_prestado']:,.2f}"],
        ["Interés mensual esperado", f"${r['total_interes_mensual']:,.2f}"],
        ["Interés pendiente acumulado", f"${r['interes_pendiente_acumulado']:,.2f}"],
        ["Interés cobrado (histórico)", f"${r['interes_cobrado_total']:,.2f}"],
        ["Préstamos activos", str(r["prestamos_activos"])],
        ["Préstamos pagados", str(r["prestamos_pagados"])],
    ]
    tabla_resumen = Table(resumen_data, colWidths=[260, 150])
    tabla_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EDF5")),
        ("TEXTCOLOR", (0, 0), (-1, -1), GRAY),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, BORDE),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
    ]))
    el.append(tabla_resumen)

    if datos["prestamos"]:
        el.append(Paragraph("Préstamos", seccion_style))
        filas = [["Fecha", "Monto", "Interés mensual", "Saldo capital", "Estado"]]
        for p in datos["prestamos"]:
            filas.append([
                p["fecha_prestamo"] or "—",
                f"${float(p['monto'] or 0):,.2f}",
                f"${float(p['interes_mensual'] or 0):,.2f}",
                f"${float(p['saldo_capital'] or 0):,.2f}",
                "Pagado" if p["pagado"] else "Activo",
            ])
        t = Table(filas, colWidths=[70, 80, 90, 90, 60])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), GOLD),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.5, BORDE),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F4F6")]),
        ]))
        el.append(t)

    pendientes = [c for c in datos["cortes"] if not c["pagado"]]
    if pendientes:
        el.append(Paragraph("Intereses pendientes de cobro", seccion_style))
        filas = [["Periodo", "Monto interés", "Tipo de pago"]]
        for c in pendientes:
            filas.append([c["periodo"] or "—", f"${float(c['monto_interes'] or 0):,.2f}", c.get("tipo_pago") or "—"])
        t = Table(filas, colWidths=[100, 100, 140])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), ORANGE),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.5, BORDE),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FEF0E3")]),
        ]))
        el.append(t)

    if datos["abonos"]:
        el.append(Paragraph("Últimos abonos de capital", seccion_style))
        filas = [["Fecha", "Interés", "Capital", "Nota"]]
        for a in datos["abonos"][:15]:
            filas.append([
                a["fecha_pago"] or "—",
                f"${float(a['monto_interes'] or 0):,.2f}",
                f"${float(a['monto_capital'] or 0):,.2f}",
                (a.get("nota") or "—")[:40],
            ])
        t = Table(filas, colWidths=[70, 80, 80, 160])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), GOLD),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.5, BORDE),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F4F6")]),
        ]))
        el.append(t)

    doc.build(el)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


@app.route("/api/clientes/<int:cliente_id>/informe-deudor/pdf", methods=["GET"])
@requiere_lectura("consultor")
def get_informe_deudor_pdf(cliente_id):
    """Descarga el informe ejecutivo del deudor como PDF con el membrete de la empresa."""
    datos = _datos_informe_deudor(cliente_id)
    pdf_bytes = _generar_pdf_informe(datos)
    archivo = f"informe_{datos['deudor'].strip().replace(' ', '_')}_{hoy_mx().isoformat()}.pdf"
    return Response(
        pdf_bytes,
        mimetype="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{archivo}"'},
    )


def _generar_xlsx_informe_deudor(datos):
    """Construye el estado de cuenta del deudor en Excel (una hoja por
       sección), a partir de los mismos datos que usa el PDF
       (_datos_informe_deudor), para que ambos formatos nunca se desalineen."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font

    negrita = Font(bold=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Estado de cuenta", datos["deudor"]])
    ws.append(["Fecha del informe", hoy_mx().isoformat()])
    ws.append([])
    r = datos["resumen"]
    for etiqueta, valor in [
        ("Capital prestado activo", r["total_prestado"]),
        ("Interés mensual esperado", r["total_interes_mensual"]),
        ("Interés pendiente acumulado", r["interes_pendiente_acumulado"]),
        ("Interés cobrado (histórico)", r["interes_cobrado_total"]),
        ("Préstamos activos", r["prestamos_activos"]),
        ("Préstamos pagados", r["prestamos_pagados"]),
    ]:
        ws.append([etiqueta, valor])
    for fila in ws.iter_rows(min_row=1, max_row=2):
        for celda in fila:
            celda.font = negrita
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("Préstamos")
    ws2.append(["Fecha", "Monto", "Interés mensual", "Saldo capital", "Estado", "Nota"])
    for celda in ws2[1]:
        celda.font = negrita
    for p in datos["prestamos"]:
        ws2.append([
            p["fecha_prestamo"] or "—",
            float(p["monto"] or 0),
            float(p["interes_mensual"] or 0),
            float(p["saldo_capital"] or 0),
            "Pagado" if p["pagado"] else "Activo",
            p.get("nota") or "",
        ])
    for col, ancho in zip("ABCDEF", (12, 14, 16, 14, 10, 40)):
        ws2.column_dimensions[col].width = ancho

    ws3 = wb.create_sheet("Intereses")
    ws3.append(["Periodo", "Monto interés", "Pagado", "Fecha de pago", "Monto pagado", "Tipo de pago"])
    for celda in ws3[1]:
        celda.font = negrita
    for co in datos["cortes"]:
        ws3.append([
            co["periodo"] or "—",
            float(co["monto_interes"] or 0),
            "Sí" if co["pagado"] else "No",
            co["fecha_pago"] or "—",
            float(co["monto_pagado"] or 0),
            co.get("tipo_pago") or "—",
        ])
    for col, ancho in zip("ABCDEF", (12, 14, 10, 14, 14, 16)):
        ws3.column_dimensions[col].width = ancho

    ws4 = wb.create_sheet("Abonos")
    ws4.append(["Fecha", "Interés", "Capital", "Tipo de pago", "Nota"])
    for celda in ws4[1]:
        celda.font = negrita
    for a in datos["abonos"]:
        ws4.append([
            a["fecha_pago"] or "—",
            float(a["monto_interes"] or 0),
            float(a["monto_capital"] or 0),
            a.get("tipo_pago") or "—",
            a.get("nota") or "",
        ])
    for col, ancho in zip("ABCDE", (12, 12, 12, 16, 40)):
        ws4.column_dimensions[col].width = ancho

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.route("/api/clientes/<int:cliente_id>/informe-deudor/xlsx", methods=["GET"])
@requiere_lectura("consultor")
def get_informe_deudor_xlsx(cliente_id):
    """Descarga el estado de cuenta del deudor en Excel (una hoja por sección: resumen, préstamos, intereses y abonos)."""
    datos = _datos_informe_deudor(cliente_id)
    xlsx_bytes = _generar_xlsx_informe_deudor(datos)
    archivo = f"estado_cuenta_{datos['deudor'].strip().replace(' ', '_')}_{hoy_mx().isoformat()}.xlsx"
    return Response(
        xlsx_bytes,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{archivo}"'},
    )


# ─── PAGARÉ / CONTRATO EN PDF POR PRÉSTAMO ────────────────────────────────────

def _generar_pdf_pagare(prestamo):
    """Construye el PDF de un pagaré simple para un préstamo. `prestamo` debe
       traer deudor_nombre, monto, interes_mensual, fecha_prestamo, telefono
       y direccion (ver la consulta en descargar_pagare)."""
    from io import BytesIO
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY

    NAVY = colors.HexColor("#0B1F4B")
    GOLD = colors.HexColor("#C9A84C")

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=18 * mm, bottomMargin=18 * mm,
                             leftMargin=22 * mm, rightMargin=22 * mm)
    styles = getSampleStyleSheet()
    titulo_style = ParagraphStyle("titulo", parent=styles["Title"], textColor=NAVY, fontSize=18, alignment=TA_CENTER, spaceAfter=2)
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], textColor=GOLD, fontSize=10, alignment=TA_CENTER, spaceAfter=18)
    cuerpo_style = ParagraphStyle("cuerpo", parent=styles["Normal"], fontSize=10.5, leading=17, spaceAfter=12, alignment=TA_JUSTIFY)
    firma_style = ParagraphStyle("firma", parent=styles["Normal"], fontSize=10, alignment=TA_CENTER)

    logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "gonza-systems.png")
    el = []
    if os.path.exists(logo_path):
        el.append(Image(logo_path, width=20 * mm, height=20 * mm, hAlign="CENTER"))
        el.append(Spacer(1, 8))

    el.append(Paragraph("PAGARÉ", titulo_style))
    el.append(Paragraph("JGM Gonzas Systems — Sistema GONZA", sub_style))

    nombre_deudor = prestamo["deudor_nombre"]
    monto = float(prestamo["monto"] or 0)
    interes = float(prestamo["interes_mensual"] or 0)
    fecha_prestamo = fmt_fecha_es(prestamo["fecha_prestamo"])
    telefono = prestamo.get("telefono") or "no registrado"
    direccion = prestamo.get("direccion") or "no registrado"

    el.append(Paragraph(
        f"Por medio del presente pagaré, yo, <b>{nombre_deudor}</b>, con teléfono {telefono} y domicilio en "
        f"{direccion}, reconozco deber y me obligo a pagar incondicionalmente a <b>JGM Gonzas Systems</b>, o a "
        f"quien sus derechos represente, la cantidad de <b>${monto:,.2f} (M.N.)</b>, recibida en calidad de "
        f"préstamo con fecha {fecha_prestamo}.",
        cuerpo_style,
    ))
    el.append(Paragraph(
        f"Sobre el saldo insoluto de esta deuda se generará un interés mensual de <b>${interes:,.2f}</b>, "
        "pagadero mes a mes en la fecha de corte correspondiente, hasta la liquidación total del capital adeudado.",
        cuerpo_style,
    ))
    el.append(Paragraph(
        "En caso de incumplimiento de pago, acepto cubrir los gastos de cobranza que se originen. "
        "Este documento es negociable conforme a la legislación aplicable en materia de títulos de crédito.",
        cuerpo_style,
    ))
    el.append(Spacer(1, 20))
    el.append(Paragraph("Lugar y fecha: ______________________________", cuerpo_style))
    el.append(Spacer(1, 40))

    firmas = Table(
        [
            ["_______________________________", "_______________________________"],
            [Paragraph("Firma del deudor", firma_style), Paragraph("Firma del acreedor", firma_style)],
            [Paragraph(nombre_deudor, firma_style), Paragraph("JGM Gonzas Systems", firma_style)],
        ],
        colWidths=[220, 220],
    )
    firmas.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 4), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    el.append(firmas)

    doc.build(el)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def fmt_fecha_es(fecha):
    if not fecha:
        return "—"
    meses = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
             "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
    texto = fecha.isoformat() if hasattr(fecha, "isoformat") else str(fecha)
    y, m, d = texto[:10].split("-")
    return f"{int(d)} de {meses[int(m)]} de {y}"


@app.route("/api/prestamos/<int:pid>/pagare", methods=["GET"])
@requiere_lectura("consultor")
def descargar_pagare(pid):
    """Genera un pagaré/contrato en PDF para un préstamo específico."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT p.deudor_nombre, p.monto, p.interes_mensual, p.fecha_prestamo::text AS fecha_prestamo,
               c.telefono, c.direccion
        FROM prestamos p
        LEFT JOIN clientes c ON c.id = p.cliente_id
        WHERE p.id = %s AND p.eliminado_en IS NULL;
    """, (pid,))
    prestamo = cur.fetchone()
    conn.close()

    if not prestamo:
        return jsonify({"error": "Préstamo no encontrado"}), 404

    pdf_bytes = _generar_pdf_pagare(prestamo)
    archivo = f"pagare_{prestamo['deudor_nombre'].replace(' ', '_')}_{pid}.pdf"
    return Response(
        pdf_bytes,
        mimetype="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{archivo}"'},
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
