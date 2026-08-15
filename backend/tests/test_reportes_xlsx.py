"""Reportes de negocio en Excel: estado de cuenta por cliente y cartera
vencida (préstamos activos con interés mensual sin cobrar)."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import load_workbook

import app as app_module


def _autenticar_como(monkeypatch, username, rol):
    monkeypatch.setattr(app_module, "_obtener_username_autenticado", lambda: (username, None))
    monkeypatch.setattr(app_module, "get_rol", lambda u: rol)


def test_estado_cuenta_xlsx_devuelve_un_excel_valido(client, monkeypatch):
    _autenticar_como(monkeypatch, "carla", "consultor")
    conn = MagicMock()
    cur = conn.cursor.return_value
    cur.fetchone.return_value = {"nombre": "Ana", "apellido_pat": "Lopez", "apellido_mat": ""}
    cur.fetchall.side_effect = [
        [{"id": 1, "deudor_nombre": "Ana Lopez", "fecha_prestamo": "2026-01-01", "monto": 1000,
          "interes_mensual": 100, "pagado": False, "fecha_pago": None, "nota": "",
          "capital_abonado": 200, "saldo_capital": 800, "tipo_pago": None}],
        [{"prestamo_id": 1, "periodo": "2026-02-01", "monto_interes": 100, "pagado": False,
          "fecha_pago": None, "monto_pagado": 0, "nota": "", "tipo_pago": None}],  # cortes_interes
        [{"prestamo_id": 1, "fecha_pago": "2026-01-15", "monto_interes": 0, "monto_capital": 200,
          "nota": "Abono", "tipo_pago": "transferencia"}],  # pagos_prestamo
    ]
    monkeypatch.setattr(app_module, "get_db", lambda: conn)

    resp = client.get("/api/clientes/7/informe-deudor/xlsx")

    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.data[:2] == b"PK"

    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Resumen", "Préstamos", "Intereses", "Abonos"]
    assert wb["Resumen"]["B1"].value == "Ana Lopez"
    assert wb["Préstamos"]["D2"].value == 800.0  # saldo capital
    assert wb["Abonos"]["C2"].value == 200.0     # monto capital abonado


def test_cartera_vencida_xlsx_devuelve_un_excel_valido(client, monkeypatch):
    _autenticar_como(monkeypatch, "carla", "consultor")
    conn = MagicMock()
    cur = conn.cursor.return_value
    cur.fetchall.side_effect = [
        [],  # préstamos activos con interés (para generar cortes faltantes) — vacío, sin efecto
        [
            {"prestamo_id": 3, "deudor_nombre": "Juan Perez", "telefono": "555-1111",
             "monto": 1000.0, "saldo": 500.0, "interes_mensual": 100.0,
             "cortes_pendientes": 2, "interes_pendiente": 200.0, "primer_corte_pendiente": "2026-01-01"},
        ],
    ]
    monkeypatch.setattr(app_module, "get_db", lambda: conn)

    resp = client.get("/api/reportes/cartera-vencida/xlsx")

    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(BytesIO(resp.data))
    ws = wb["Cartera vencida"]
    assert ws["B2"].value == "Juan Perez"
    assert ws["E2"].value == 500.0   # saldo
    assert ws["H2"].value == 200.0   # interés pendiente
    assert ws["E4"].value == 500.0   # fila de totales (saldo)
    assert ws["H4"].value == 200.0   # fila de totales (interés pendiente)


def test_estado_cuenta_xlsx_requiere_sesion(client, monkeypatch):
    monkeypatch.setattr(app_module, "_obtener_username_autenticado", lambda: (None, "No autorizado"))

    resp = client.get("/api/clientes/7/informe-deudor/xlsx")

    assert resp.status_code == 401
